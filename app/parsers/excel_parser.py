from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.cell.cell import Cell
from pathlib import Path
from typing import Dict, List, Any, Optional, Tuple
from dataclasses import dataclass, field
from collections import defaultdict
import html


@dataclass
class CellStyle:
    font: Dict[str, Any] = field(default_factory=dict)
    fill: Dict[str, Any] = field(default_factory=dict)
    border: Dict[str, Any] = field(default_factory=dict)
    alignment: Dict[str, Any] = field(default_factory=dict)
    is_merged: bool = False
    merge_range: Optional[Tuple[int, int, int, int]] = None


@dataclass
class ExcelCell:
    row: int
    col: int
    value: Any
    data_type: str
    style: CellStyle
    is_merged: bool = False
    merge_master: bool = False
    merge_range: Optional[Tuple[int, int, int, int]] = None


@dataclass
class ExcelSheet:
    name: str
    rows: List[List[ExcelCell]]
    merged_cells: List[Dict[str, Any]]
    max_row: int
    max_col: int
    column_widths: Dict[int, float]
    row_heights: Dict[int, float]


@dataclass
class ParseResult:
    success: bool
    error_message: str = ""
    sheets: List[ExcelSheet] = field(default_factory=list)
    text_content: str = ""
    html_content: str = ""
    raw_data: Dict[str, Any] = field(default_factory=dict)


def _get_color_value(color_obj) -> Optional[str]:
    """
    安全地获取颜色值，处理RGB对象和字符串
    """
    if color_obj is None:
        return None
    if hasattr(color_obj, 'rgb'):
        rgb = color_obj.rgb
        if rgb is not None:
            if isinstance(rgb, str):
                return rgb
            if hasattr(rgb, '__str__'):
                return str(rgb)
    return None


def extract_cell_style(cell: Cell) -> CellStyle:
    style = CellStyle()
    
    if cell.font:
        font = cell.font
        style.font = {
            'name': font.name,
            'size': font.size,
            'bold': font.bold,
            'italic': font.italic,
            'underline': font.underline,
            'strike': font.strike,
            'color': _get_color_value(font.color),
        }
    
    if cell.fill:
        fill = cell.fill
        style.fill = {
            'pattern_type': fill.patternType,
            'fg_color': _get_color_value(fill.fgColor),
            'bg_color': _get_color_value(fill.bgColor),
        }
    
    if cell.border:
        border = cell.border
        style.border = {
            'left': {
                'style': border.left.style if border.left else None,
                'color': _get_color_value(border.left.color) if border.left else None,
            } if border.left else None,
            'right': {
                'style': border.right.style if border.right else None,
                'color': _get_color_value(border.right.color) if border.right else None,
            } if border.right else None,
            'top': {
                'style': border.top.style if border.top else None,
                'color': _get_color_value(border.top.color) if border.top else None,
            } if border.top else None,
            'bottom': {
                'style': border.bottom.style if border.bottom else None,
                'color': _get_color_value(border.bottom.color) if border.bottom else None,
            } if border.bottom else None,
        }
    
    if cell.alignment:
        alignment = cell.alignment
        style.alignment = {
            'horizontal': alignment.horizontal,
            'vertical': alignment.vertical,
            'text_rotation': alignment.textRotation,
            'wrap_text': alignment.wrapText,
            'shrink_to_fit': alignment.shrinkToFit,
            'indent': alignment.indent,
        }
    
    return style


def get_data_type(cell: Cell) -> str:
    if cell.data_type == 'n':
        if cell.number_format and ('%' in cell.number_format or 'percentage' in cell.number_format.lower()):
            return 'percentage'
        if cell.number_format and any(x in cell.number_format.lower() for x in ['date', 'time', 'yyyy', 'mm', 'dd']):
            return 'datetime'
        return 'number'
    elif cell.data_type == 's':
        return 'string'
    elif cell.data_type == 'b':
        return 'boolean'
    elif cell.data_type == 'd':
        return 'datetime'
    elif cell.data_type == 'f':
        return 'formula'
    elif cell.data_type == 'n' and cell.value is None:
        return 'empty'
    else:
        return 'other'


def parse_excel_file(file_path: str) -> ParseResult:
    """
    解析Excel文件，返回详细的解析结果
    """
    file_path = Path(file_path)
    
    try:
        wb = load_workbook(file_path, data_only=False, read_only=False)
    except Exception as e:
        return ParseResult(success=False, error_message=f"无法加载Excel文件: {str(e)}")
    
    sheets = []
    all_text_parts = []
    all_html_parts = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        merged_cells_info = []
        merged_cell_ranges = {}
        
        for merged_range in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged_range.bounds
            merged_info = {
                'min_row': min_row,
                'min_col': min_col,
                'max_row': max_row,
                'max_col': max_col,
                'range_string': str(merged_range),
            }
            merged_cells_info.append(merged_info)
            
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    merged_cell_ranges[(row, col)] = {
                        'is_master': (row == min_row and col == min_col),
                        'min_row': min_row,
                        'min_col': min_col,
                        'max_row': max_row,
                        'max_col': max_col,
                    }
        
        max_row = ws.max_row if ws.max_row else 0
        max_col = ws.max_column if ws.max_column else 0
        
        column_widths = {}
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            column_widths[col_idx] = ws.column_dimensions.get(col_letter, None).width if ws.column_dimensions.get(col_letter) else None
        
        row_heights = {}
        for row_idx in range(1, max_row + 1):
            row_heights[row_idx] = ws.row_dimensions.get(row_idx, None).height if ws.row_dimensions.get(row_idx) else None
        
        rows_data = []
        sheet_text_parts = []
        
        for row_idx in range(1, max_row + 1):
            row_cells = []
            row_text_values = []
            
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                style = extract_cell_style(cell)
                
                is_merged = False
                merge_master = False
                merge_range = None
                
                if (row_idx, col_idx) in merged_cell_ranges:
                    merge_info = merged_cell_ranges[(row_idx, col_idx)]
                    is_merged = True
                    merge_master = merge_info['is_master']
                    merge_range = (
                        merge_info['min_row'],
                        merge_info['min_col'],
                        merge_info['max_row'],
                        merge_info['max_col'],
                    )
                    style.is_merged = True
                    style.merge_range = merge_range
                
                data_type = get_data_type(cell)
                
                cell_value = cell.value
                if cell.data_type == 'f' and cell_value is not None:
                    try:
                        if cell.value.startswith('='):
                            pass
                    except:
                        pass
                
                excel_cell = ExcelCell(
                    row=row_idx,
                    col=col_idx,
                    value=cell_value,
                    data_type=data_type,
                    style=style,
                    is_merged=is_merged,
                    merge_master=merge_master,
                    merge_range=merge_range,
                )
                
                row_cells.append(excel_cell)
                
                if cell_value is not None:
                    cell_text = str(cell_value)
                    row_text_values.append(cell_text)
            
            rows_data.append(row_cells)
            
            if row_text_values:
                sheet_text_parts.append(" | ".join(row_text_values))
        
        excel_sheet = ExcelSheet(
            name=sheet_name,
            rows=rows_data,
            merged_cells=merged_cells_info,
            max_row=max_row,
            max_col=max_col,
            column_widths=column_widths,
            row_heights=row_heights,
        )
        sheets.append(excel_sheet)
        
        all_text_parts.append(f"=== 工作表: {sheet_name} ===")
        all_text_parts.extend(sheet_text_parts)
        
        html_table = convert_sheet_to_html(excel_sheet)
        all_html_parts.append(f"<h2>工作表: {sheet_name}</h2>")
        all_html_parts.append(html_table)
    
    text_content = "\n\n".join(all_text_parts)
    html_content = "\n".join(all_html_parts)
    
    wb.close()
    
    return ParseResult(
        success=True,
        sheets=sheets,
        text_content=text_content,
        html_content=html_content,
        raw_data={
            'sheet_count': len(sheets),
            'sheet_names': [s.name for s in sheets],
        }
    )


def convert_sheet_to_html(sheet: ExcelSheet) -> str:
    """
    将Excel工作表转换为HTML表格
    """
    html_parts = []
    
    style_parts = [
        "border-collapse: collapse;",
        "width: 100%;",
        "margin: 10px 0;",
        "font-family: Arial, sans-serif;",
        "font-size: 14px;",
    ]
    html_parts.append(f"<table style='{''.join(style_parts)}'>")
    
    merged_covered = set()
    
    for row_idx, row in enumerate(sheet.rows, 1):
        html_parts.append("<tr>")
        
        for col_idx, cell in enumerate(row, 1):
            if (row_idx, col_idx) in merged_covered:
                continue
            
            cell_style = []
            cell_attrs = []
            
            if cell.is_merged and cell.merge_master and cell.merge_range:
                min_row, min_col, max_row, max_col = cell.merge_range
                rowspan = max_row - min_row + 1
                colspan = max_col - min_col + 1
                
                if rowspan > 1:
                    cell_attrs.append(f"rowspan='{rowspan}'")
                if colspan > 1:
                    cell_attrs.append(f"colspan='{colspan}'")
                
                for r in range(min_row, max_row + 1):
                    for c in range(min_col, max_col + 1):
                        merged_covered.add((r, c))
            
            if cell.style.font:
                font = cell.style.font
                if font.get('bold'):
                    cell_style.append("font-weight: bold;")
                if font.get('italic'):
                    cell_style.append("font-style: italic;")
                if font.get('underline') and font.get('underline') != 'none':
                    cell_style.append("text-decoration: underline;")
                if font.get('color'):
                    color = font.get('color')
                    color_str = str(color) if color is not None else ""
                    if len(color_str) >= 6:
                        if color_str.startswith('FF') and len(color_str) > 2:
                            color_str = color_str[2:]
                        if len(color_str) >= 6:
                            cell_style.append(f"color: #{color_str};")
                if font.get('size'):
                    cell_style.append(f"font-size: {font.get('size')}pt;")
            
            if cell.style.fill:
                fill = cell.style.fill
                if fill.get('pattern_type') and fill.get('pattern_type') != 'none':
                    fg_color = fill.get('fg_color')
                    fg_color_str = str(fg_color) if fg_color is not None else ""
                    if len(fg_color_str) >= 6:
                        if fg_color_str.startswith('FF') and len(fg_color_str) > 2:
                            fg_color_str = fg_color_str[2:]
                        if len(fg_color_str) >= 6:
                            cell_style.append(f"background-color: #{fg_color_str};")
            
            if cell.style.alignment:
                alignment = cell.style.alignment
                if alignment.get('horizontal'):
                    h_align = alignment.get('horizontal')
                    align_map = {
                        'left': 'left',
                        'center': 'center',
                        'right': 'right',
                        'justify': 'justify',
                        'general': 'left',
                    }
                    if h_align in align_map:
                        cell_style.append(f"text-align: {align_map[h_align]};")
                if alignment.get('vertical'):
                    v_align = alignment.get('vertical')
                    valign_map = {
                        'top': 'top',
                        'center': 'middle',
                        'bottom': 'bottom',
                    }
                    if v_align in valign_map:
                        cell_style.append(f"vertical-align: {valign_map[v_align]};")
                if alignment.get('wrap_text'):
                    cell_style.append("white-space: normal;")
                    cell_style.append("word-wrap: break-word;")
            
            cell_style.append("border: 1px solid #ccc;")
            cell_style.append("padding: 8px;")
            
            tag = "th" if row_idx == 1 else "td"
            
            style_attr = f" style='{''.join(cell_style)}'" if cell_style else ""
            attrs_str = f" {' '.join(cell_attrs)}" if cell_attrs else ""
            
            cell_value = cell.value
            if cell_value is None:
                cell_text = ""
            else:
                cell_text = html.escape(str(cell_value))
            
            html_parts.append(f"<{tag}{style_attr}{attrs_str}>{cell_text}</{tag}>")
        
        html_parts.append("</tr>")
    
    html_parts.append("</table>")
    
    return "\n".join(html_parts)


def parse_excel(file_path: str) -> tuple[str, str]:
    """
    兼容旧接口的解析函数，返回原始内容和HTML内容
    """
    result = parse_excel_file(file_path)
    if result.success:
        return result.text_content, result.html_content
    else:
        return "", ""


def update_cell_in_excel(file_path: str, sheet_name: str, row: int, col: int, value: Any) -> bool:
    """
    更新Excel文件中指定单元格的值
    
    Args:
        file_path: Excel文件路径
        sheet_name: 工作表名称
        row: 行号（从1开始）
        col: 列号（从1开始）
        value: 新值
    
    Returns:
        是否成功
    """
    try:
        wb = load_workbook(file_path)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return False
        
        ws = wb[sheet_name]
        cell = ws.cell(row=row, column=col)
        cell.value = value
        wb.save(file_path)
        wb.close()
        return True
    except Exception as e:
        print(f"更新单元格失败: {e}")
        return False


def update_multiple_cells(file_path: str, updates: List[Dict[str, Any]]) -> bool:
    """
    批量更新Excel文件中的多个单元格
    
    Args:
        file_path: Excel文件路径
        updates: 更新列表，每个元素包含:
            - sheet_name: 工作表名称
            - row: 行号
            - col: 列号
            - value: 新值
    
    Returns:
        是否成功
    """
    try:
        wb = load_workbook(file_path)
        
        for update in updates:
            sheet_name = update.get('sheet_name')
            row = update.get('row')
            col = update.get('col')
            value = update.get('value')
            
            if sheet_name not in wb.sheetnames:
                continue
            
            ws = wb[sheet_name]
            cell = ws.cell(row=row, column=col)
            cell.value = value
        
        wb.save(file_path)
        wb.close()
        return True
    except Exception as e:
        print(f"批量更新单元格失败: {e}")
        return False


def add_new_sheet(file_path: str, sheet_name: str) -> bool:
    """
    在Excel文件中添加新工作表
    
    Args:
        file_path: Excel文件路径
        sheet_name: 新工作表名称
    
    Returns:
        是否成功
    """
    try:
        wb = load_workbook(file_path)
        
        if sheet_name in wb.sheetnames:
            wb.close()
            return False
        
        wb.create_sheet(sheet_name)
        wb.save(file_path)
        wb.close()
        return True
    except Exception as e:
        print(f"添加工作表失败: {e}")
        return False


def delete_sheet(file_path: str, sheet_name: str) -> bool:
    """
    删除Excel文件中的指定工作表
    
    Args:
        file_path: Excel文件路径
        sheet_name: 要删除的工作表名称
    
    Returns:
        是否成功
    """
    try:
        wb = load_workbook(file_path)
        
        if sheet_name not in wb.sheetnames:
            wb.close()
            return False
        
        if len(wb.sheetnames) <= 1:
            wb.close()
            return False
        
        del wb[sheet_name]
        wb.save(file_path)
        wb.close()
        return True
    except Exception as e:
        print(f"删除工作表失败: {e}")
        return False


def get_sheet_data_as_json(file_path: str, sheet_name: Optional[str] = None) -> Dict[str, Any]:
    """
    获取Excel文件数据为JSON格式
    
    Args:
        file_path: Excel文件路径
        sheet_name: 工作表名称（可选，默认获取所有工作表）
    
    Returns:
        包含数据的字典
    """
    result = parse_excel_file(file_path)
    
    if not result.success:
        return {'success': False, 'error': result.error_message}
    
    sheets_data = []
    
    for sheet in result.sheets:
        if sheet_name and sheet.name != sheet_name:
            continue
        
        sheet_dict = {
            'name': sheet.name,
            'max_row': sheet.max_row,
            'max_col': sheet.max_col,
            'merged_cells': sheet.merged_cells,
            'data': []
        }
        
        for row in sheet.rows:
            row_data = []
            for cell in row:
                cell_dict = {
                    'row': cell.row,
                    'col': cell.col,
                    'value': cell.value,
                    'data_type': cell.data_type,
                    'is_merged': cell.is_merged,
                    'merge_master': cell.merge_master,
                }
                if cell.merge_range:
                    cell_dict['merge_range'] = cell.merge_range
                row_data.append(cell_dict)
            sheet_dict['data'].append(row_data)
        
        sheets_data.append(sheet_dict)
    
    return {
        'success': True,
        'sheets': sheets_data,
        'sheet_count': len(sheets_data),
    }
