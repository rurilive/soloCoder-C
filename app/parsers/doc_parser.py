from pathlib import Path
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from zipfile import BadZipFile
import olefile
import struct
import re
import chardet
import logging
from logging.handlers import RotatingFileHandler
import os
from dataclasses import dataclass, field
from typing import Optional, List, Dict, Any, Tuple, Callable
from abc import ABC, abstractmethod
from collections import Counter
import html


LOG_DIR = Path(__file__).resolve().parent.parent.parent / 'log'
LOG_DIR.mkdir(parents=True, exist_ok=True)
LOG_FILE = LOG_DIR / 'doc_parser.log'

logger = logging.getLogger('doc_parser')
logger.setLevel(logging.DEBUG)

formatter = logging.Formatter('%(asctime)s - %(name)s - %(levelname)s - %(message)s')

file_handler = RotatingFileHandler(
    LOG_FILE, 
    maxBytes=10*1024*1024,
    backupCount=5,
    encoding='utf-8'
)
file_handler.setLevel(logging.DEBUG)
file_handler.setFormatter(formatter)

console_handler = logging.StreamHandler()
console_handler.setLevel(logging.INFO)
console_handler.setFormatter(formatter)

if not logger.handlers:
    logger.addHandler(file_handler)
    logger.addHandler(console_handler)

logger.info("=" * 60)
logger.info("doc_parser 模块初始化")
logger.info(f"日志文件: {LOG_FILE}")
logger.info("=" * 60)


try:
    from pyantiword.antiword_wrapper import extract_text_with_antiword as antiword_extract_text
    HAS_ANTIWORD = True
    logger.info("✅ pyantiword 已导入，函数: extract_text_with_antiword")
except ImportError as e:
    HAS_ANTIWORD = False
    logger.warning(f"⚠️ pyantiword 导入失败: {e}")

try:
    import mammoth
    HAS_MAMMOTH = True
    logger.info("✅ mammoth 已导入，可以更好地解析 docx 格式")
except ImportError as e:
    HAS_MAMMOTH = False
    logger.warning(f"⚠️ mammoth 导入失败: {e}，将使用 python-docx 作为备选")

try:
    from cmi_docx import Document as CmiDocument
    HAS_CMI_DOCX = True
    logger.info("✅ cmi-docx 已导入，提供增强的 docx 编辑和解析功能")
except ImportError as e:
    HAS_CMI_DOCX = False
    logger.warning(f"⚠️ cmi-docx 导入失败: {e}")


@dataclass
class ParseResultMetrics:
    total_text_length: int = 0
    paragraph_count: int = 0
    heading_count: int = 0
    heading_levels: Dict[int, int] = field(default_factory=dict)
    list_item_count: int = 0
    ordered_list_count: int = 0
    unordered_list_count: int = 0
    table_count: int = 0
    table_row_count: int = 0
    table_cell_count: int = 0
    inline_format_count: int = 0
    bold_count: int = 0
    italic_count: int = 0
    underline_count: int = 0
    link_count: int = 0
    image_count: int = 0
    style_attributes_count: int = 0
    alignment_count: int = 0
    indent_count: int = 0
    spacing_count: int = 0
    semantic_tag_ratio: float = 0.0
    non_empty_paragraph_ratio: float = 0.0


@dataclass
class ParseResult:
    parser_name: str
    text_content: str = ""
    html_content: str = ""
    success: bool = True
    error_message: str = ""
    metrics: ParseResultMetrics = field(default_factory=ParseResultMetrics)
    raw_data: Dict[str, Any] = field(default_factory=dict)
    score: float = 0.0


class BaseDocxParser(ABC):
    @abstractmethod
    def parse(self, file_path: str) -> ParseResult:
        pass
    
    @property
    @abstractmethod
    def name(self) -> str:
        pass
    
    @property
    @abstractmethod
    def priority(self) -> int:
        pass


class MetricsCalculator:
    @staticmethod
    def calculate_metrics(text_content: str, html_content: str, raw_data: Dict = None) -> ParseResultMetrics:
        metrics = ParseResultMetrics()
        metrics.total_text_length = len(text_content) if text_content else 0
        
        if html_content:
            metrics = MetricsCalculator._analyze_html(html_content, metrics)
        
        if raw_data:
            metrics = MetricsCalculator._enrich_from_raw_data(raw_data, metrics)
        
        metrics = MetricsCalculator._calculate_ratios(metrics, text_content)
        
        return metrics
    
    @staticmethod
    def _analyze_html(html_content: str, metrics: ParseResultMetrics) -> ParseResultMetrics:
        soup = None
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html_content, 'html.parser')
        except ImportError:
            pass
        
        if soup:
            return MetricsCalculator._analyze_with_bs4(soup, metrics)
        else:
            return MetricsCalculator._analyze_with_regex(html_content, metrics)
    
    @staticmethod
    def _analyze_with_bs4(soup, metrics: ParseResultMetrics) -> ParseResultMetrics:
        paragraphs = soup.find_all(['p', 'div'])
        metrics.paragraph_count = len(paragraphs)
        
        non_empty_paras = sum(1 for p in paragraphs if p.get_text(strip=True))
        if metrics.paragraph_count > 0:
            metrics.non_empty_paragraph_ratio = non_empty_paras / metrics.paragraph_count
        
        for level in range(1, 7):
            headings = soup.find_all(f'h{level}')
            count = len(headings)
            if count > 0:
                metrics.heading_levels[level] = count
                metrics.heading_count += count
        
        lists = soup.find_all(['ul', 'ol'])
        for lst in lists:
            items = lst.find_all('li')
            metrics.list_item_count += len(items)
            if lst.name == 'ol':
                metrics.ordered_list_count += 1
            else:
                metrics.unordered_list_count += 1
        
        tables = soup.find_all('table')
        metrics.table_count = len(tables)
        for table in tables:
            rows = table.find_all('tr')
            metrics.table_row_count += len(rows)
            for row in rows:
                cells = row.find_all(['td', 'th'])
                metrics.table_cell_count += len(cells)
        
        bold_tags = soup.find_all(['strong', 'b'])
        metrics.bold_count = len(bold_tags)
        italic_tags = soup.find_all(['em', 'i'])
        metrics.italic_count = len(italic_tags)
        underline_tags = soup.find_all(['u'])
        metrics.underline_count = len(underline_tags)
        metrics.inline_format_count = metrics.bold_count + metrics.italic_count + metrics.underline_count
        
        links = soup.find_all('a', href=True)
        metrics.link_count = len(links)
        
        images = soup.find_all('img')
        metrics.image_count = len(images)
        
        for tag in soup.find_all(True):
            if tag.has_attr('style'):
                style = tag['style']
                metrics.style_attributes_count += 1
                if 'text-align' in style:
                    metrics.alignment_count += 1
                if 'text-indent' in style or 'margin-left' in style:
                    metrics.indent_count += 1
                if 'margin-top' in style or 'margin-bottom' in style:
                    metrics.spacing_count += 1
        
        return metrics
    
    @staticmethod
    def _analyze_with_regex(html_content: str, metrics: ParseResultMetrics) -> ParseResultMetrics:
        metrics.paragraph_count = len(re.findall(r'<p\b[^>]*>', html_content, re.IGNORECASE))
        
        for level in range(1, 7):
            pattern = rf'<h{level}\b[^>]*>'
            count = len(re.findall(pattern, html_content, re.IGNORECASE))
            if count > 0:
                metrics.heading_levels[level] = count
                metrics.heading_count += count
        
        ol_count = len(re.findall(r'<ol\b[^>]*>', html_content, re.IGNORECASE))
        ul_count = len(re.findall(r'<ul\b[^>]*>', html_content, re.IGNORECASE))
        metrics.ordered_list_count = ol_count
        metrics.unordered_list_count = ul_count
        metrics.list_item_count = len(re.findall(r'<li\b[^>]*>', html_content, re.IGNORECASE))
        
        metrics.table_count = len(re.findall(r'<table\b[^>]*>', html_content, re.IGNORECASE))
        metrics.table_row_count = len(re.findall(r'<tr\b[^>]*>', html_content, re.IGNORECASE))
        metrics.table_cell_count = len(re.findall(r'<td\b[^>]*>|<th\b[^>]*>', html_content, re.IGNORECASE))
        
        metrics.bold_count = len(re.findall(r'<strong\b[^>]*>|<b\b[^>]*>', html_content, re.IGNORECASE))
        metrics.italic_count = len(re.findall(r'<em\b[^>]*>|<i\b[^>]*>', html_content, re.IGNORECASE))
        metrics.underline_count = len(re.findall(r'<u\b[^>]*>', html_content, re.IGNORECASE))
        metrics.inline_format_count = metrics.bold_count + metrics.italic_count + metrics.underline_count
        
        metrics.link_count = len(re.findall(r'<a\b[^>]*href=[\'"]', html_content, re.IGNORECASE))
        metrics.image_count = len(re.findall(r'<img\b[^>]*>', html_content, re.IGNORECASE))
        
        style_matches = re.findall(r'style=[\'"]([^\'"]*)[\'"]', html_content, re.IGNORECASE)
        metrics.style_attributes_count = len(style_matches)
        for style in style_matches:
            if 'text-align' in style:
                metrics.alignment_count += 1
            if 'text-indent' in style or 'margin-left' in style:
                metrics.indent_count += 1
            if 'margin-top' in style or 'margin-bottom' in style:
                metrics.spacing_count += 1
        
        return metrics
    
    @staticmethod
    def _enrich_from_raw_data(raw_data: Dict, metrics: ParseResultMetrics) -> ParseResultMetrics:
        if 'paragraphs_info' in raw_data:
            paras = raw_data['paragraphs_info']
            metrics.paragraph_count = max(metrics.paragraph_count, len(paras))
            
            for para in paras:
                if para.get('styles'):
                    metrics.style_attributes_count += 1
                    style = para['styles']
                    if 'text-align' in style:
                        metrics.alignment_count += 1
                    if 'text-indent' in style:
                        metrics.indent_count += 1
                    if 'margin' in style:
                        metrics.spacing_count += 1
        
        if 'tables' in raw_data:
            tables = raw_data['tables']
            metrics.table_count = max(metrics.table_count, len(tables))
            for table in tables:
                rows = table.get('rows', [])
                metrics.table_row_count += len(rows)
                for row in rows:
                    metrics.table_cell_count += len(row)
        
        return metrics
    
    @staticmethod
    def _calculate_ratios(metrics: ParseResultMetrics, text_content: str) -> ParseResultMetrics:
        total_structural = (
            metrics.heading_count + 
            metrics.list_item_count + 
            metrics.table_count +
            metrics.inline_format_count
        )
        
        if metrics.paragraph_count > 0:
            metrics.semantic_tag_ratio = total_structural / metrics.paragraph_count
        
        return metrics


class ResultScorer:
    WEIGHTS = {
        'text_length': 0.10,
        'heading_count': 0.15,
        'list_item_count': 0.10,
        'table_count': 0.10,
        'table_cell_count': 0.05,
        'inline_format_count': 0.10,
        'link_count': 0.05,
        'image_count': 0.05,
        'style_attributes_count': 0.15,
        'alignment_count': 0.05,
        'indent_count': 0.05,
        'semantic_tag_ratio': 0.05,
    }
    
    @staticmethod
    def calculate_score(result: ParseResult, all_results: List[ParseResult] = None) -> float:
        metrics = result.metrics
        score = 0.0
        
        max_values = ResultScorer._get_max_values(all_results) if all_results else None
        
        score += ResultScorer._normalize(metrics.total_text_length, max_values.get('text_length') if max_values else None, 10000) * ResultScorer.WEIGHTS['text_length']
        
        score += ResultScorer._normalize(metrics.heading_count, max_values.get('heading_count') if max_values else None, 20) * ResultScorer.WEIGHTS['heading_count']
        
        score += ResultScorer._normalize(metrics.list_item_count, max_values.get('list_item_count') if max_values else None, 50) * ResultScorer.WEIGHTS['list_item_count']
        
        score += ResultScorer._normalize(metrics.table_count, max_values.get('table_count') if max_values else None, 10) * ResultScorer.WEIGHTS['table_count']
        
        score += ResultScorer._normalize(metrics.table_cell_count, max_values.get('table_cell_count') if max_values else None, 100) * ResultScorer.WEIGHTS['table_cell_count']
        
        score += ResultScorer._normalize(metrics.inline_format_count, max_values.get('inline_format_count') if max_values else None, 100) * ResultScorer.WEIGHTS['inline_format_count']
        
        score += ResultScorer._normalize(metrics.link_count, max_values.get('link_count') if max_values else None, 20) * ResultScorer.WEIGHTS['link_count']
        
        score += ResultScorer._normalize(metrics.image_count, max_values.get('image_count') if max_values else None, 10) * ResultScorer.WEIGHTS['image_count']
        
        score += ResultScorer._normalize(metrics.style_attributes_count, max_values.get('style_attributes_count') if max_values else None, 50) * ResultScorer.WEIGHTS['style_attributes_count']
        
        score += ResultScorer._normalize(metrics.alignment_count, max_values.get('alignment_count') if max_values else None, 20) * ResultScorer.WEIGHTS['alignment_count']
        
        score += ResultScorer._normalize(metrics.indent_count, max_values.get('indent_count') if max_values else None, 20) * ResultScorer.WEIGHTS['indent_count']
        
        score += min(metrics.semantic_tag_ratio, 1.0) * ResultScorer.WEIGHTS['semantic_tag_ratio']
        
        return min(score, 1.0)
    
    @staticmethod
    def _get_max_values(results: List[ParseResult]) -> Dict[str, float]:
        max_vals = {
            'text_length': 0,
            'heading_count': 0,
            'list_item_count': 0,
            'table_count': 0,
            'table_cell_count': 0,
            'inline_format_count': 0,
            'link_count': 0,
            'image_count': 0,
            'style_attributes_count': 0,
            'alignment_count': 0,
            'indent_count': 0,
        }
        
        for result in results:
            m = result.metrics
            max_vals['text_length'] = max(max_vals['text_length'], m.total_text_length)
            max_vals['heading_count'] = max(max_vals['heading_count'], m.heading_count)
            max_vals['list_item_count'] = max(max_vals['list_item_count'], m.list_item_count)
            max_vals['table_count'] = max(max_vals['table_count'], m.table_count)
            max_vals['table_cell_count'] = max(max_vals['table_cell_count'], m.table_cell_count)
            max_vals['inline_format_count'] = max(max_vals['inline_format_count'], m.inline_format_count)
            max_vals['link_count'] = max(max_vals['link_count'], m.link_count)
            max_vals['image_count'] = max(max_vals['image_count'], m.image_count)
            max_vals['style_attributes_count'] = max(max_vals['style_attributes_count'], m.style_attributes_count)
            max_vals['alignment_count'] = max(max_vals['alignment_count'], m.alignment_count)
            max_vals['indent_count'] = max(max_vals['indent_count'], m.indent_count)
        
        return max_vals
    
    @staticmethod
    def _normalize(value: float, max_value: float = None, fallback_max: float = 100) -> float:
        if max_value is None or max_value == 0:
            max_value = fallback_max
        if max_value == 0:
            return 0.0
        return min(value / max_value, 1.0)


class PythonDocxParser(BaseDocxParser):
    @property
    def name(self) -> str:
        return "python-docx"
    
    @property
    def priority(self) -> int:
        return 3
    
    def parse(self, file_path: str) -> ParseResult:
        logger.info(f"[PythonDocxParser] 开始解析: {file_path}")
        
        try:
            doc = DocxDocument(file_path)
            
            content_parts = []
            html_parts = []
            paragraphs_info = []
            tables_info = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                content_parts.append(text)
                
                style_name = para.style.name.lower() if para.style.name else ''
                full_styles = self._get_paragraph_styles(para)
                
                paragraphs_info.append({
                    'text': text,
                    'styles': full_styles,
                    'alignment': str(para.alignment) if para.alignment else None,
                    'style_name': para.style.name,
                })
                
                style_attr = f' style="{full_styles}"' if full_styles else ''
                
                if 'heading 1' in style_name:
                    html_parts.append(f"<h1{style_attr}>{self._escape_html(text)}</h1>")
                elif 'heading 2' in style_name:
                    html_parts.append(f"<h2{style_attr}>{self._escape_html(text)}</h2>")
                elif 'heading 3' in style_name:
                    html_parts.append(f"<h3{style_attr}>{self._escape_html(text)}</h3>")
                elif 'heading 4' in style_name:
                    html_parts.append(f"<h4{style_attr}>{self._escape_html(text)}</h4>")
                elif 'heading 5' in style_name:
                    html_parts.append(f"<h5{style_attr}>{self._escape_html(text)}</h5>")
                elif 'heading 6' in style_name:
                    html_parts.append(f"<h6{style_attr}>{self._escape_html(text)}</h6>")
                elif 'list' in style_name or 'bullet' in style_name:
                    html_parts.append(f"<li{style_attr}>{self._escape_html(text)}</li>")
                else:
                    html_parts.append(f"<p{style_attr}>{self._escape_html(text)}</p>")
            
            for table in doc.tables:
                table_data = {'rows': []}
                html_table = "<table border='1' cellpadding='5' cellspacing='0' style='border-collapse: collapse;'>"
                for row_idx, row in enumerate(table.rows):
                    row_data = []
                    html_table += "<tr>"
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        row_data.append(cell_text)
                        cell_tag = 'th' if row_idx == 0 else 'td'
                        html_table += f"<{cell_tag}>{self._escape_html(cell_text)}</{cell_tag}>"
                    html_table += "</tr>"
                    table_data['rows'].append(row_data)
                    content_parts.append(" | ".join(row_data))
                html_table += "</table>"
                html_parts.append(html_table)
                tables_info.append(table_data)
            
            text_content = "\n\n".join(content_parts)
            html_content = "\n".join(html_parts)
            
            metrics = MetricsCalculator.calculate_metrics(
                text_content, 
                html_content,
                {'paragraphs_info': paragraphs_info, 'tables': tables_info}
            )
            
            result = ParseResult(
                parser_name=self.name,
                text_content=text_content,
                html_content=html_content,
                success=True,
                metrics=metrics,
                raw_data={
                    'paragraphs_info': paragraphs_info,
                    'tables': tables_info,
                    'doc_object': doc,
                }
            )
            
            logger.info(f"[PythonDocxParser] 解析完成，文本长度: {len(text_content)}, HTML长度: {len(html_content)}")
            return result
            
        except Exception as e:
            logger.error(f"[PythonDocxParser] 解析失败: {e}")
            import traceback
            traceback.print_exc()
            return ParseResult(
                parser_name=self.name,
                success=False,
                error_message=str(e)
            )
    
    def _get_paragraph_styles(self, para) -> str:
        styles = []
        
        align_style = self._get_alignment_style(para.alignment)
        if align_style:
            styles.append(align_style)
        
        pf = para.paragraph_format
        
        first_line_indent = pf.first_line_indent
        if first_line_indent is not None:
            try:
                indent_pt = first_line_indent.pt
                if indent_pt > 0:
                    indent_em = indent_pt / 12
                    styles.append(f"text-indent: {indent_em:.2f}em;")
            except:
                pass
        
        left_indent = pf.left_indent
        if left_indent is not None:
            try:
                indent_pt = left_indent.pt
                if indent_pt > 0:
                    indent_em = indent_pt / 12
                    styles.append(f"margin-left: {indent_em:.2f}em;")
            except:
                pass
        
        right_indent = pf.right_indent
        if right_indent is not None:
            try:
                indent_pt = right_indent.pt
                if indent_pt > 0:
                    indent_em = indent_pt / 12
                    styles.append(f"margin-right: {indent_em:.2f}em;")
            except:
                pass
        
        space_before = pf.space_before
        if space_before is not None:
            try:
                pt = space_before.pt
                if pt > 0:
                    styles.append(f"margin-top: {pt:.1f}pt;")
            except:
                pass
        
        space_after = pf.space_after
        if space_after is not None:
            try:
                pt = space_after.pt
                if pt > 0:
                    styles.append(f"margin-bottom: {pt:.1f}pt;")
            except:
                pass
        
        return " ".join(styles)
    
    def _get_alignment_style(self, alignment) -> str:
        if alignment is None:
            return ""
        
        alignment_map = {
            WD_ALIGN_PARAGRAPH.LEFT: "",
            WD_ALIGN_PARAGRAPH.CENTER: "text-align: center;",
            WD_ALIGN_PARAGRAPH.RIGHT: "text-align: right;",
            WD_ALIGN_PARAGRAPH.JUSTIFY: "text-align: justify;",
            WD_ALIGN_PARAGRAPH.DISTRIBUTE: "text-align: justify; text-align-last: justify;",
            WD_ALIGN_PARAGRAPH.JUSTIFY_MED: "text-align: justify;",
            WD_ALIGN_PARAGRAPH.JUSTIFY_HI: "text-align: justify;",
            WD_ALIGN_PARAGRAPH.JUSTIFY_LOW: "text-align: justify;",
            WD_ALIGN_PARAGRAPH.THAI_JUSTIFY: "text-align: justify;",
        }
        
        return alignment_map.get(alignment, "")
    
    def _escape_html(self, text: str) -> str:
        return html.escape(text)


class MammothParser(BaseDocxParser):
    @property
    def name(self) -> str:
        return "mammoth"
    
    @property
    def priority(self) -> int:
        return 2
    
    def parse(self, file_path: str) -> ParseResult:
        logger.info(f"[MammothParser] 开始解析: {file_path}")
        
        if not HAS_MAMMOTH:
            return ParseResult(
                parser_name=self.name,
                success=False,
                error_message="mammoth 库未安装"
            )
        
        try:
            with open(file_path, "rb") as docx_file:
                result = mammoth.convert_to_html(docx_file)
                html_content = result.value
                messages = result.messages
            
            if messages:
                logger.warning(f"[MammothParser] 转换消息: {messages}")
            
            paragraphs_info = self._extract_paragraphs_from_docx(file_path)
            html_content = self._enhance_html_with_alignment(html_content, file_path, paragraphs_info)
            
            text_content = re.sub(r'<[^>]+>', ' ', html_content)
            text_content = re.sub(r'\s+', ' ', text_content).strip()
            
            metrics = MetricsCalculator.calculate_metrics(
                text_content,
                html_content,
                {'paragraphs_info': paragraphs_info}
            )
            
            result = ParseResult(
                parser_name=self.name,
                text_content=text_content,
                html_content=html_content,
                success=True,
                metrics=metrics,
                raw_data={
                    'paragraphs_info': paragraphs_info,
                    'messages': messages,
                }
            )
            
            logger.info(f"[MammothParser] 解析完成，文本长度: {len(text_content)}, HTML长度: {len(html_content)}")
            return result
            
        except Exception as e:
            logger.error(f"[MammothParser] 解析失败: {e}")
            import traceback
            traceback.print_exc()
            return ParseResult(
                parser_name=self.name,
                success=False,
                error_message=str(e)
            )
    
    def _extract_paragraphs_from_docx(self, file_path: str) -> List[Dict]:
        try:
            doc = DocxDocument(file_path)
            paragraphs_info = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    styles = self._get_paragraph_styles(para)
                    paragraphs_info.append({
                        'text': text,
                        'styles': styles,
                        'alignment': para.alignment,
                        'indent': para.paragraph_format.first_line_indent,
                        'used': False
                    })
            return paragraphs_info
        except:
            return []
    
    def _get_paragraph_styles(self, para) -> str:
        styles = []
        
        align_style = self._get_alignment_style(para.alignment)
        if align_style:
            styles.append(align_style)
        
        pf = para.paragraph_format
        
        first_line_indent = pf.first_line_indent
        if first_line_indent is not None:
            try:
                indent_pt = first_line_indent.pt
                if indent_pt > 0:
                    indent_em = indent_pt / 12
                    styles.append(f"text-indent: {indent_em:.2f}em;")
            except:
                pass
        
        left_indent = pf.left_indent
        if left_indent is not None:
            try:
                indent_pt = left_indent.pt
                if indent_pt > 0:
                    indent_em = indent_pt / 12
                    styles.append(f"margin-left: {indent_em:.2f}em;")
            except:
                pass
        
        return " ".join(styles)
    
    def _get_alignment_style(self, alignment) -> str:
        if alignment is None:
            return ""
        
        alignment_map = {
            WD_ALIGN_PARAGRAPH.LEFT: "",
            WD_ALIGN_PARAGRAPH.CENTER: "text-align: center;",
            WD_ALIGN_PARAGRAPH.RIGHT: "text-align: right;",
            WD_ALIGN_PARAGRAPH.JUSTIFY: "text-align: justify;",
        }
        
        return alignment_map.get(alignment, "")
    
    def _enhance_html_with_alignment(self, html_content: str, file_path: str, paragraphs_info: List[Dict]) -> str:
        if not paragraphs_info:
            return html_content
        
        def add_styles_to_tag(match):
            tag = match.group(1)
            existing_attrs = match.group(2) or ''
            content = match.group(3)
            
            content_clean = re.sub(r'<[^>]+>', '', content).strip()
            
            unused_paragraphs = [p for p in paragraphs_info if not p['used']]
            para_info = self._find_best_match(content_clean, unused_paragraphs)
            
            if para_info:
                para_info['used'] = True
                if para_info['styles']:
                    if 'style=' in existing_attrs:
                        existing_attrs = re.sub(
                            r'style="([^"]*)"',
                            f'style="\\1 {para_info["styles"]}"',
                            existing_attrs
                        )
                    else:
                        existing_attrs = f' style="{para_info["styles"]}" {existing_attrs}'
            
            return f'<{tag}{existing_attrs.strip()}>{content}</{tag}>'
        
        html_content = re.sub(
            r'<(h[1-6]|p|li)([^>]*)>(.*?)</\1>',
            add_styles_to_tag,
            html_content,
            flags=re.DOTALL | re.IGNORECASE
        )
        
        return html_content
    
    def _find_best_match(self, html_text: str, paragraphs_info: List[Dict]) -> Optional[Dict]:
        if not html_text:
            return None
        
        html_normalized = self._normalize_text(html_text)
        if not html_normalized:
            return None
        
        best_match = None
        best_score = 0
        
        for para_info in paragraphs_info:
            para_text = para_info.get('text', '')
            para_normalized = self._normalize_text(para_text)
            
            if not para_normalized:
                continue
            
            score = 0
            if html_normalized == para_normalized:
                score = 100
            elif para_normalized in html_normalized or html_normalized in para_normalized:
                score = 50
            else:
                min_len = min(len(html_normalized), len(para_normalized))
                if min_len > 0:
                    common_chars = sum(1 for c in html_normalized if c in para_normalized)
                    score = (common_chars / min_len) * 30
            
            if score > best_score and score > 10:
                best_score = score
                best_match = para_info
        
        return best_match
    
    def _normalize_text(self, text: str) -> str:
        if not text:
            return ""
        text = re.sub(r'\s+', '', text)
        text = re.sub(r'[，。、；：""\(\)【】《》！？\.\,\;\:\'\"\[\]\<\>\!\?]', '', text)
        return text.lower()


class CmiDocxParser(BaseDocxParser):
    @property
    def name(self) -> str:
        return "cmi-docx"
    
    @property
    def priority(self) -> int:
        return 1
    
    def parse(self, file_path: str) -> ParseResult:
        logger.info(f"[CmiDocxParser] 开始解析: {file_path}")
        
        if not HAS_CMI_DOCX:
            return ParseResult(
                parser_name=self.name,
                success=False,
                error_message="cmi-docx 库未安装"
            )
        
        try:
            doc = CmiDocument(file_path)
            
            content_parts = []
            html_parts = []
            paragraphs_info = []
            
            for para_idx, para in enumerate(doc.paragraphs):
                text = para.text.strip() if hasattr(para, 'text') else str(para).strip()
                if not text:
                    continue
                
                content_parts.append(text)
                
                style_info = {}
                if hasattr(para, 'style'):
                    style_info['style_name'] = para.style
                
                full_styles = ""
                if hasattr(para, 'alignment'):
                    align_style = self._get_alignment_style(para.alignment)
                    if align_style:
                        full_styles = align_style
                
                paragraphs_info.append({
                    'text': text,
                    'styles': full_styles,
                    'style_info': style_info,
                    'index': para_idx,
                })
                
                style_attr = f' style="{full_styles}"' if full_styles else ''
                
                style_name = style_info.get('style_name', '').lower() if style_info.get('style_name') else ''
                
                if 'heading 1' in style_name or 'h1' in style_name:
                    html_parts.append(f"<h1{style_attr}>{self._escape_html(text)}</h1>")
                elif 'heading 2' in style_name or 'h2' in style_name:
                    html_parts.append(f"<h2{style_attr}>{self._escape_html(text)}</h2>")
                elif 'heading 3' in style_name or 'h3' in style_name:
                    html_parts.append(f"<h3{style_attr}>{self._escape_html(text)}</h3>")
                elif 'heading' in style_name:
                    html_parts.append(f"<h4{style_attr}>{self._escape_html(text)}</h4>")
                else:
                    html_parts.append(f"<p{style_attr}>{self._escape_html(text)}</p>")
            
            tables_info = []
            if hasattr(doc, 'tables'):
                for table_idx, table in enumerate(doc.tables):
                    table_data = {'rows': []}
                    html_table = "<table border='1' cellpadding='5' cellspacing='0' style='border-collapse: collapse;'>"
                    
                    if hasattr(table, 'rows'):
                        for row_idx, row in enumerate(table.rows):
                            row_data = []
                            html_table += "<tr>"
                            
                            if hasattr(row, 'cells'):
                                for cell in row.cells:
                                    cell_text = cell.text.strip() if hasattr(cell, 'text') else str(cell).strip()
                                    row_data.append(cell_text)
                                    cell_tag = 'th' if row_idx == 0 else 'td'
                                    html_table += f"<{cell_tag}>{self._escape_html(cell_text)}</{cell_tag}>"
                            
                            html_table += "</tr>"
                            table_data['rows'].append(row_data)
                            if row_data:
                                content_parts.append(" | ".join(row_data))
                    
                    html_table += "</table>"
                    html_parts.append(html_table)
                    tables_info.append(table_data)
            
            text_content = "\n\n".join(content_parts)
            html_content = "\n".join(html_parts)
            
            metrics = MetricsCalculator.calculate_metrics(
                text_content,
                html_content,
                {'paragraphs_info': paragraphs_info, 'tables': tables_info}
            )
            
            result = ParseResult(
                parser_name=self.name,
                text_content=text_content,
                html_content=html_content,
                success=True,
                metrics=metrics,
                raw_data={
                    'paragraphs_info': paragraphs_info,
                    'tables': tables_info,
                    'cmi_document': doc,
                }
            )
            
            logger.info(f"[CmiDocxParser] 解析完成，文本长度: {len(text_content)}, HTML长度: {len(html_content)}")
            return result
            
        except Exception as e:
            logger.error(f"[CmiDocxParser] 解析失败: {e}")
            import traceback
            traceback.print_exc()
            return ParseResult(
                parser_name=self.name,
                success=False,
                error_message=str(e)
            )
    
    def _get_alignment_style(self, alignment) -> str:
        if alignment is None:
            return ""
        
        align_str = str(alignment).lower()
        
        if 'center' in align_str:
            return "text-align: center;"
        elif 'right' in align_str:
            return "text-align: right;"
        elif 'justify' in align_str:
            return "text-align: justify;"
        
        return ""
    
    def _escape_html(self, text: str) -> str:
        return html.escape(text)


class MultiParserCompetitor:
    def __init__(self):
        self._parsers: List[BaseDocxParser] = []
        self._register_parsers()
    
    def _register_parsers(self):
        self._parsers.append(PythonDocxParser())
        self._parsers.append(MammothParser())
        if HAS_CMI_DOCX:
            self._parsers.append(CmiDocxParser())
        
        self._parsers.sort(key=lambda p: p.priority, reverse=True)
        
        logger.info(f"[MultiParserCompetitor] 已注册 {len(self._parsers)} 个解析器: {[p.name for p in self._parsers]}")
    
    def parse_all(self, file_path: str) -> List[ParseResult]:
        logger.info(f"[MultiParserCompetitor] 开始多解析器竞争解析: {file_path}")
        
        results = []
        
        for parser in self._parsers:
            logger.info(f"[MultiParserCompetitor] 调用解析器: {parser.name}")
            result = parser.parse(file_path)
            results.append(result)
            
            if result.success:
                logger.info(f"[MultiParserCompetitor] 解析器 {parser.name} 成功")
            else:
                logger.warning(f"[MultiParserCompetitor] 解析器 {parser.name} 失败: {result.error_message}")
        
        successful_results = [r for r in results if r.success]
        logger.info(f"[MultiParserCompetitor] 成功解析器数量: {len(successful_results)}/{len(results)}")
        
        if successful_results:
            for result in successful_results:
                result.score = ResultScorer.calculate_score(result, successful_results)
            
            successful_results.sort(key=lambda r: r.score, reverse=True)
            
            logger.info("[MultiParserCompetitor] 解析器评分排名:")
            for idx, result in enumerate(successful_results):
                logger.info(f"  {idx + 1}. {result.parser_name}: 得分={result.score:.4f}")
                m = result.metrics
                logger.info(f"     - 文本长度: {m.total_text_length}, 标题数: {m.heading_count}")
                logger.info(f"     - 列表项: {m.list_item_count}, 表格数: {m.table_count}")
                logger.info(f"     - 样式属性: {m.style_attributes_count}, 内联格式: {m.inline_format_count}")
        
        return results
    
    def select_best(self, results: List[ParseResult]) -> Optional[ParseResult]:
        successful_results = [r for r in results if r.success]
        
        if not successful_results:
            logger.error("[MultiParserCompetitor] 没有成功的解析结果")
            return None
        
        successful_results.sort(key=lambda r: r.score, reverse=True)
        
        best = successful_results[0]
        logger.info(f"[MultiParserCompetitor] 选择最优解析器: {best.parser_name}, 得分: {best.score:.4f}")
        
        return best
    
    def parse_and_select(self, file_path: str) -> Tuple[Optional[ParseResult], List[ParseResult]]:
        results = self.parse_all(file_path)
        best = self.select_best(results)
        return best, results


_multi_parser_competitor: Optional[MultiParserCompetitor] = None


def get_multi_parser_competitor() -> MultiParserCompetitor:
    global _multi_parser_competitor
    if _multi_parser_competitor is None:
        _multi_parser_competitor = MultiParserCompetitor()
    return _multi_parser_competitor


def is_docx_file(file_path: str) -> bool:
    logger.debug(f"[is_docx_file] 检查文件: {file_path}")
    with open(file_path, 'rb') as f:
        header = f.read(4)
        result = header == b'PK\x03\x04'
        logger.debug(f"[is_docx_file] 文件头: {header}, 结果: {result}")
        return result


def is_doc_file(file_path: str) -> bool:
    logger.debug(f"[is_doc_file] 检查文件: {file_path}")
    result = olefile.isOleFile(file_path)
    logger.debug(f"[is_doc_file] 结果: {result}")
    return result


def _is_unicode_word(word_data: bytes) -> bool:
    try:
        flags = struct.unpack('<H', word_data[0xA:0xC])[0]
        return (flags & 0x0100) != 0
    except:
        return False


def _detect_encoding(data: bytes) -> str:
    if not data or len(data) == 0:
        return 'utf-8'
    
    try:
        result = chardet.detect(data)
        if result and result['encoding']:
            confidence = result.get('confidence', 0)
            logger.debug(f"[_detect_encoding] 检测到编码: {result['encoding']}, 置信度: {confidence}")
            if confidence > 0.5:
                return result['encoding']
    except:
        pass
    
    return 'utf-8'


def _smart_clean_text(text: str) -> str:
    if not text:
        return ""
    
    logger.debug(f"[_smart_clean_text] 清理前长度: {len(text)}")
    
    lines = text.split('\n')
    cleaned_lines = []
    
    for line in lines:
        line = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F]', '', line)
        line = line.rstrip()
        cleaned_lines.append(line)
    
    result = '\n'.join(cleaned_lines)
    
    logger.debug(f"[_smart_clean_text] 清理后长度: {len(result)}")
    
    return result


def _is_highly_likely_valid_text(text: str) -> bool:
    if not text or len(text.strip()) < 20:
        return False
    
    total_chars = len(text)
    if total_chars == 0:
        return False
    
    printable_chars = sum(1 for c in text if c.isprintable() or c in '\n\r\t ')
    printable_ratio = printable_chars / total_chars
    
    if printable_ratio < 0.7:
        return False
    
    chinese_chars = sum(1 for c in text if '\u4e00' <= c <= '\u9fff')
    english_letters = sum(1 for c in text if c.isalpha() and ord(c) < 128)
    digits = sum(1 for c in text if c.isdigit())
    
    meaningful_chars = chinese_chars + english_letters + digits
    meaningful_ratio = meaningful_chars / total_chars
    
    if meaningful_ratio < 0.3:
        return False
    
    weird_chars = sum(1 for c in text if (
        '\u0080' <= c <= '\u00FF' or
        '\u0100' <= c <= '\u017F' or
        '\u0180' <= c <= '\u024F' or
        '\u0250' <= c <= '\u02AF' or
        '\u1E00' <= c <= '\u1EFF' or
        '\u2C60' <= c <= '\u2C7F' or
        '\uA720' <= c <= '\uA7FF' or
        '\uAB30' <= c <= '\uAB6F' or
        '\uFB00' <= c <= '\uFB4F' or
        '\uFF00' <= c <= '\uFFEF' or
        c in '⸀㬆ᨆఀ䭓쌀ꐀ؆㔊᐀Ȁ␀㨀☆ᔀ䤊꬀ᔊ Ȁఁ䘀'
    ))
    
    if weird_chars > total_chars * 0.1:
        return False
    
    control_chars = sum(1 for c in text if (
        '\u2000' <= c <= '\u206F' or
        '\uFFF0' <= c <= '\uFFFF' or
        '\uE000' <= c <= '\uF8FF'
    ))
    
    if control_chars > total_chars * 0.05:
        return False
    
    if chinese_chars > 10:
        return True
    
    if english_letters > 50:
        common_words = ['the', 'and', 'for', 'are', 'but', 'not', 'you', 'all', 'can', 'has', 'her', 'was', 'one', 'our', 'out', 'day', 'get', 'has', 'him', 'his', 'how', 'its', 'may', 'new', 'now', 'old', 'see', 'two', 'way', 'who', 'boy', 'did', 'own', 'say', 'she', 'too', 'use', 'with', 'have', 'this', 'will', 'your', 'from', 'they', 'been', 'call', 'make', 'come', 'time', 'just', 'know', 'take', 'into', 'year', 'some', 'want', 'look', 'work', 'give', 'over', 'think', 'most', 'find', 'day', 'also', 'after', 'way', 'many', 'must', 'look', 'before', 'great', 'back', 'through', 'long', 'where', 'much', 'should', 'well', 'people', 'down', 'own', 'upon', 'good', 'part', 'place', 'write', 'word', 'about', 'read', 'man', 'find', 'change', 'went', 'light', 'kind', 'off', 'need', 'house', 'picture', 'try', 'again', 'animal', 'point', 'mother', 'world', 'near', 'build', 'self', 'earth', 'father', 'head', 'stand', 'page', 'country', 'found', 'answer', 'school', 'grow', 'study', 'learn', 'plant', 'cover', 'food', 'sun', 'thought', 'let', 'keep', 'eye', 'never', 'last', 'door', 'between', 'city', 'tree', 'cross', 'since', 'hard', 'start', 'might', 'story', 'saw', 'far', 'sea', 'draw', 'left', 'late', 'run', "don't", 'while', 'press', 'close', 'night', 'real', 'life', 'few', 'north', 'book', 'carry', 'took', 'science', 'eat', 'room', 'friend', 'began', 'idea', 'fish', 'mountain', 'stop', 'once', 'base', 'hear', 'horse', 'cut', 'sure', 'watch', 'color', 'face', 'wood', 'main', 'open', 'seem', 'together', 'next', 'white', 'children', 'begin', 'got', 'walk', 'example', 'ease', 'paper', 'group', 'always', 'music', 'those', 'both', 'mark', 'often', 'letter', 'until', 'mile', 'river', 'car', 'feet', 'care', 'second', 'book', 'remember', 'early', 'game', 'line', 'quite', 'move', 'thing', 'light', 'kilogram', 'common', 'come', 'difference', 'use', 'language', 'govern', 'please', 'simple', 'third', 'order', 'fire', 'south', 'problem', 'full', 'hear', 'ask', 'force', 'air', 'today', 'important', 'man', 'women', 'child', 'family', 'home', 'work', 'study', 'school', 'student', 'teacher', 'doctor', 'nurse', 'office', 'company', 'business', 'market', 'price', 'money', 'bank', 'loan', 'payment', 'cost', 'expense', 'income', 'salary', 'wage', 'tax', 'government', 'law', 'policy', 'rule', 'right', 'duty', 'responsibility', 'free', 'choice', 'decision', 'plan', 'action', 'step', 'process', 'system', 'method', 'way', 'approach', 'technique', 'skill', 'ability', 'talent', 'gift', 'power', 'strength', 'energy', 'force', 'pressure', 'stress', 'pain', 'pleasure', 'happiness', 'sadness', 'anger', 'fear', 'love', 'hate', 'like', 'dislike', 'want', 'need', 'desire', 'hope', 'dream', 'goal', 'aim', 'purpose', 'reason', 'cause', 'effect', 'result', 'change', 'development', 'growth', 'progress', 'success', 'failure', 'win', 'lose', 'gain', 'loss', 'advantage', 'disadvantage', 'benefit', 'harm', 'good', 'bad', 'better', 'worse', 'best', 'worst', 'more', 'less', 'most', 'least', 'many', 'few', 'much', 'little', 'some', 'any', 'all', 'none', 'each', 'every', 'both', 'either', 'neither', 'same', 'different', 'similar', 'like', 'unlike', 'equal', 'unequal', 'big', 'small', 'large', 'little', 'long', 'short', 'wide', 'narrow', 'high', 'low', 'deep', 'shallow', 'thick', 'thin', 'heavy', 'light', 'fast', 'slow', 'quick', 'slow', 'early', 'late', 'old', 'young', 'new', 'first', 'last', 'beginning', 'end', 'start', 'finish', 'open', 'close', 'on', 'off', 'in', 'out', 'up', 'down', 'left', 'right', 'front', 'back', 'top', 'bottom', 'above', 'below', 'under', 'over', 'between', 'among', 'through', 'across', 'along', 'around', 'about', 'before', 'after', 'during', 'while', 'since', 'until', 'when', 'where', 'why', 'how', 'what', 'which', 'who', 'whom', 'whose', 'this', 'that', 'these', 'those', 'here', 'there', 'now', 'then', 'today', 'tomorrow', 'yesterday', 'week', 'month', 'year', 'time', 'hour', 'minute', 'second', 'day', 'night', 'morning', 'afternoon', 'evening', 'spring', 'summer', 'autumn', 'winter', 'season', 'weather', 'rain', 'sun', 'wind', 'cloud', 'snow', 'ice', 'fire', 'water', 'air', 'earth', 'wood', 'metal', 'stone', 'sand', 'dust', 'mud', 'blood', 'bone', 'skin', 'hair', 'eye', 'ear', 'nose', 'mouth', 'hand', 'foot', 'leg', 'arm', 'head', 'body', 'heart', 'mind', 'soul', 'spirit', 'thought', 'idea', 'feeling', 'emotion', 'memory', 'imagination', 'creativity', 'intelligence', 'wisdom', 'knowledge', 'information', 'data', 'fact', 'truth', 'lie', 'fiction', 'story', 'history', 'news', 'report', 'article', 'book', 'magazine', 'newspaper', 'journal', 'diary', 'note', 'letter', 'message', 'email', 'phone', 'computer', 'internet', 'web', 'site', 'page', 'link', 'network', 'server', 'client', 'software', 'hardware', 'program', 'code', 'algorithm', 'function', 'method', 'class', 'object', 'variable', 'constant', 'parameter', 'argument', 'return', 'value', 'type', 'string', 'number', 'integer', 'float', 'boolean', 'list', 'array', 'dict', 'map', 'set', 'tuple', 'record', 'struct', 'class', 'interface', 'module', 'package', 'library', 'framework', 'api', 'database', 'table', 'row', 'column', 'query', 'sql', 'transaction', 'commit', 'rollback', 'index', 'key', 'primary', 'foreign', 'unique', 'constraint', 'trigger', 'procedure', 'function', 'view', 'cursor', 'lock', 'transaction', 'isolation', 'level', 'read', 'write', 'update', 'delete', 'insert', 'select', 'from', 'where', 'join', 'inner', 'outer', 'left', 'right', 'full', 'cross', 'natural', 'union', 'intersect', 'except', 'group', 'order', 'having', 'limit', 'offset', 'fetch', 'into', 'as', 'distinct', 'all', 'top', 'percent', 'with', 'ties', 'case', 'when', 'then', 'else', 'end', 'coalesce', 'nullif', 'isnull', 'ifnull', 'nvl', 'decode', 'greatest', 'least', 'abs', 'ceil', 'floor', 'round', 'trunc', 'mod', 'power', 'sqrt', 'exp', 'ln', 'log', 'log10', 'sin', 'cos', 'tan', 'asin', 'acos', 'atan', 'sinh', 'cosh', 'tanh', 'sign', 'chr', 'ascii', 'length', 'substr', 'substring', 'instr', 'locate', 'position', 'concat', 'upper', 'lower', 'initcap', 'trim', 'ltrim', 'rtrim', 'lpad', 'rpad', 'replace', 'translate', 'regexp_replace', 'regexp_substr', 'regexp_instr', 'regexp_like', 'to_char', 'to_number', 'to_date', 'current_date', 'current_time', 'current_timestamp', 'sysdate', 'now', 'today', 'yesterday', 'tomorrow', 'add_months', 'months_between', 'last_day', 'next_day', 'trunc', 'round', 'extract', 'date_part', 'date_trunc', 'age', 'overlaps', 'at', 'time', 'zone', 'local', 'session', 'user', 'current_user', 'session_user', 'system_user', 'public', 'role', 'grant', 'revoke', 'privilege', 'permission', 'allow', 'deny', 'enable', 'disable', 'create', 'alter', 'drop', 'truncate', 'comment', 'rename', 'add', 'modify', 'change', 'column', 'constraint', 'index', 'view', 'sequence', 'synonym', 'database', 'schema', 'table', 'tablespace', 'datafile', 'logfile', 'controlfile', 'parameter', 'spfile', 'pfile', 'instance', 'session', 'process', 'thread', 'memory', 'cache', 'buffer', 'pool', 'sga', 'pga', 'uga', 'shared', 'large', 'java', 'streams', 'background', 'foreground', 'smon', 'pmon', 'dbwr', 'lgwr', 'ckpt', 'arcn', 'reco', 'cjq0', 'qmn', 'dmon', 'smon', 'pmon', 'mmon', 'mmnl', 'fvfm', 'lmhb', 'lmd', 'lmon', 'lms', 'lmns', 'diag', 'dia0', 'dia1', 'dbw', 'lgw', 'ckp', 'arc', 'rman', 'exp', 'imp', 'expdp', 'impdp', 'sqlplus', 'sqlldr', 'tnslsnr', 'lsnrctl', 'agent', 'oem', 'grid', 'cluster', 'rac', 'dg', 'dataguard', 'standby', 'primary', 'logical', 'physical', 'snapshot', 'materialized', 'view', 'partition', 'subpartition', 'range', 'hash', 'list', 'composite', 'interval', 'reference', 'system', 'local', 'global', 'index', 'bitmap', 'function-based', 'domain', 'reverse', 'descending', 'ascending', 'unique', 'non-unique', 'b-tree', 'hash', 'cluster', 'iot', 'organization', 'heap', 'temporary', 'permanent', 'undo', 'rollback', 'temp', 'default', 'temporary', 'tablespace', 'bigfile', 'smallfile', 'autoextend', 'maxsize', 'uniform', 'autoallocate', 'segment', 'extent', 'block', 'size', 'initial', 'next', 'minextents', 'maxextents', 'pctincrease', 'freelists', 'freelist', 'groups', 'buffer', 'pool', 'keep', 'recycle', 'default', 'flash_cache', 'cell', 'flash_cache', 'none', 'default', 'keep', 'none', 'read', 'write', 'read write', 'read only', 'offline', 'online', 'begin', 'backup', 'end', 'backup', 'force', 'logging', 'nologging', 'force', 'logging', 'minimize', 'records', 'per_block', 'check', 'checksum', 'none', 'db_block_checking', 'db_block_checksum', 'db_lost_write_protect', 'parameter', 'file', 'system', 'statistics_level', 'timed_statistics', 'statistics_level', 'all', 'typical', 'basic', 'memory_target', 'memory_max_target', 'sga_target', 'sga_max_size', 'pga_aggregate_target', 'pga_aggregate_limit', 'large_pool_size', 'java_pool_size', 'streams_pool_size', 'shared_pool_size', 'db_cache_size', 'db_keep_cache_size', 'db_recycle_cache_size', 'db_nk_cache_size', 'log_buffer', 'processes', 'sessions', 'transactions', 'open_cursors', 'session_cached_cursors', 'cursor_space_for_time', 'cursor_sharing', 'exact', 'similar', 'force', 'optimizer_mode', 'all_rows', 'first_rows', 'first_rows_n', 'rule', 'choose', 'optimizer_index_caching', 'optimizer_index_cost_adj', 'db_file_multiblock_read_count', 'db_file_name_convert', 'log_file_name_convert', 'control_files', 'db_files', 'db_name', 'db_unique_name', 'instance_name', 'service_names', 'db_domain', 'global_names', 'remote_login_passwordfile', 'exclusive', 'shared', 'none', 'utl_file_dir', 'audit_file_dest', 'background_dump_dest', 'user_dump_dest', 'core_dump_dest', 'diagnostic_dest', 'compatible', 'plsql_code_type', 'native', 'interpreted', 'plsql_optimize_level', 'plsql_warnings', 'nls_language', 'nls_territory', 'nls_characterset', 'nls_nchar_characterset', 'nls_sort', 'nls_comp', 'nls_date_format', 'nls_date_language', 'nls_calendar', 'nls_numeric_characters', 'nls_currency', 'nls_iso_currency', 'nls_dual_currency', 'nls_sort', 'nls_comp', 'nls_length_semantics', 'byte', 'char', 'time_zone', 'sessiontimezone', 'dbtimezone', 'instance', 'session', 'user', 'sys', 'system', 'sysman', 'dbsnmp', 'outln', 'perfstat', 'statspack', 'mdsys', 'sdo', 'ordsys', 'ordim', 'si', 'ctxsys', 'xdb', 'sysaux', 'users', 'example', 'temp', 'undotbs1', 'undotbs2', 'system', 'sysaux', 'users', 'tools', 'indx', 'drsys', 'xdb', 'olap', 'cwmlite', 'oem_repository', 'management', 'repository', 'grid', 'cluster', 'asm', 'spfile', 'pfile', 'init', 'orapw', 'password', 'file', 'listener', 'tnsnames', 'sqlnet', 'protocol', 'address', 'host', 'port', 'service_name', 'sid', 'connect_data', 'server', 'dedicated', 'shared', 'pooled', 'dispatchers', 'max_dispatchers', 'shared_servers', 'max_shared_servers', 'large_pool_size', 'java_pool_size', 'streams_pool_size', 'job_queue_processes', 'aq_tm_processes', 'db_writer_processes', 'log_writer', 'checkpoint', 'archiver', 'reco', 'smon', 'pmon', 'mmon', 'mmnl', 'diag', 'dia0', 'dia1', 'fmom', 'fmon', 'gmon', 'imr', 'lmd', 'lmon', 'lms', 'lmns', 'lock', 'lck0', 'lck1', 'mark', 'ktm', 'mmon', 'mmnl', 'mman', 'psp0', 'qmn', 'qmn0', 'qmn1', 'reco', 'smco', 'smon', 'vkrm', 'vktm', 'w000', 'w001', 'wmon', 'z000', 'z001']
        
        text_lower = text.lower()
        for word in common_words:
            if f' {word} ' in text_lower or text_lower.startswith(f'{word} ') or text_lower.endswith(f' {word}'):
                return True
    
    if text.count(' ') > len(text) * 0.03 and text.count('\n') > 0:
        return True
    
    return False


def _is_likely_valid_text(text: str) -> bool:
    if not text or len(text.strip()) < 10:
        return False
    
    total_chars = len(text)
    if total_chars == 0:
        return False
    
    printable_chars = sum(1 for c in text if c.isprintable() or c in '\n\r\t ')
    printable_ratio = printable_chars / total_chars
    
    if printable_ratio < 0.6:
        return False
    
    ascii_chars = sum(1 for c in text if ord(c) < 128)
    ascii_ratio = ascii_chars / total_chars
    
    if ascii_ratio > 0.3:
        common_words = ['the', 'and', 'for', 'are', 'but', 'not', 'you', 'all', 'can', 'has', 'her', 'was', 'one', 'our', 'out', 'day', 'get', 'has', 'him', 'his', 'how', 'its', 'may', 'new', 'now', 'old', 'see', 'two', 'way', 'who', 'boy', 'did', 'own', 'say', 'she', 'too', 'use']
        text_lower = text.lower()
        for word in common_words:
            if f' {word} ' in text_lower or text_lower.startswith(f'{word} ') or text_lower.endswith(f' {word}'):
                return True
    
    chinese_chars = sum(1 for c in text if '\u4e00' <= c <= '\u9fff')
    if chinese_chars > 5:
        return True
    
    if text.count(' ') > len(text) * 0.05 and text.count('\n') > 0:
        return True
    
    return False


def _smart_split_paragraphs(text: str) -> list:
    if not text:
        return []
    
    logger.debug(f"[_smart_split_paragraphs] 原始文本长度: {len(text)}")
    
    lines = text.split('\n')
    logger.debug(f"[_smart_split_paragraphs] 按 \\n 分割为 {len(lines)} 行")
    
    paragraphs = []
    current_para = []
    blank_line_count = 0
    
    for i, line in enumerate(lines):
        line = line.rstrip()
        
        if not line:
            blank_line_count += 1
            if current_para:
                para_text = ' '.join(current_para).strip()
                if para_text:
                    paragraphs.append(para_text)
                current_para = []
            continue
        
        blank_line_count = 0
        
        line_stripped = line.strip()
        
        if line.startswith('|') and line.endswith('|'):
            if current_para:
                para_text = ' '.join(current_para).strip()
                if para_text:
                    paragraphs.append(para_text)
                current_para = []
            paragraphs.append(line)
            continue
        
        leading_spaces = len(line) - len(line.lstrip())
        
        if leading_spaces > 10 and current_para:
            logger.debug(f"[_smart_split_paragraphs] 行 {i} 有大量前导空格 ({leading_spaces}个)，可能是新段落: {repr(line_stripped[:30])}")
            para_text = ' '.join(current_para).strip()
            if para_text:
                paragraphs.append(para_text)
            current_para = [line_stripped]
            continue
        
        if current_para:
            last_para = current_para[-1]
            if last_para and last_para.endswith(('。', '！', '？', '；', '：', '.', '!', '?', ';', ':')):
                logger.debug(f"[_smart_split_paragraphs] 前一行以标点结尾，可能是新段落")
                para_text = ' '.join(current_para).strip()
                if para_text:
                    paragraphs.append(para_text)
                current_para = [line_stripped]
                continue
        
        current_para.append(line_stripped)
    
    if current_para:
        para_text = ' '.join(current_para).strip()
        if para_text:
            paragraphs.append(para_text)
    
    logger.debug(f"[_smart_split_paragraphs] 最终分割为 {len(paragraphs)} 个段落")
    for i, para in enumerate(paragraphs):
        logger.debug(f"  [{i}] {repr(para[:50])}")
    
    return paragraphs


def _is_likely_english_heading(para: str) -> tuple:
    if not para or len(para.strip()) == 0:
        return False, None
    
    para = para.strip()
    para_len = len(para)
    
    if para_len > 80:
        return False, None
    
    chinese_count = sum(1 for c in para if '\u4e00' <= c <= '\u9fff')
    english_letters = sum(1 for c in para if c.isalpha() and ord(c) < 128)
    
    total_chars = chinese_count + english_letters
    if total_chars == 0:
        return False, None
    
    english_ratio = english_letters / total_chars
    logger.debug(f"[_is_likely_english_heading] 英文比例: {english_ratio:.2f}, 长度: {para_len}, 内容: {repr(para)}")
    
    if chinese_count == 0 and english_ratio > 0.8:
        logger.debug(f"[_is_likely_english_heading] 纯英文段落，可能是标题")
        
        if para.isupper() and para_len >= 2:
            logger.debug(f"[_is_likely_english_heading] 全部大写英文，识别为标题")
            if para_len <= 10:
                return True, 'h1'
            elif para_len <= 20:
                return True, 'h2'
            else:
                return True, 'h3'
        
        words = para.split()
        if len(words) >= 1 and len(words) <= 10:
            title_case = all(word and word[0].isupper() for word in words if word)
            if title_case and not para.endswith(('.', '!', '?', ',', ';', ':')):
                logger.debug(f"[_is_likely_english_heading] 标题格式英文，识别为标题")
                if para_len <= 15:
                    return True, 'h2'
                elif para_len <= 30:
                    return True, 'h3'
                elif para_len <= 50:
                    return True, 'h4'
                else:
                    return True, 'h5'
        
        if english_ratio > 0.9 and para_len <= 30:
            if not para.endswith(('.', '!', '?', ',', ';', ':')):
                logger.debug(f"[_is_likely_english_heading] 纯英文短句，识别为标题")
                if para_len <= 15:
                    return True, 'h3'
                else:
                    return True, 'h4'
    
    return False, None


def _is_likely_heading(para: str, prev_para: str = None, next_para: str = None) -> tuple:
    if not para or len(para.strip()) == 0:
        return False, None
    
    para = para.strip()
    para_len = len(para)
    
    if para.startswith('|') and para.endswith('|'):
        return False, None
    
    heading_patterns = [
        (r'^第[一二三四五六七八九十百千\d]+[章节篇条]\s*[：:]*\s*', 1),
        (r'^[一二三四五六七八九十]+[、.．]\s*', 2),
        (r'^(\d+)[\.．、]\s*', 2),
        (r'^(\d+\.\d+)[\.．\s]*', 3),
        (r'^(\d+\.\d+\.\d+)[\.．\s]*', 4),
        (r'^[（(]\d+[)）]\s*', 3),
        (r'^[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]\s*', 3),
    ]
    
    for pattern, level in heading_patterns:
        if re.match(pattern, para):
            logger.debug(f"[_is_likely_heading] 匹配标题模式: {para[:30]}..., 级别: h{level}")
            return True, f'h{level}'
    
    is_eng_heading, eng_level = _is_likely_english_heading(para)
    if is_eng_heading:
        logger.debug(f"[_is_likely_heading] 识别为英文标题: {para[:30]}..., 级别: {eng_level}")
        return True, eng_level
    
    if para_len <= 80:
        ends_with_punctuation = para.endswith(('。', '！', '？', '，', '；', '：', '.', '!', '?', ',', ';', ':'))
        
        if not ends_with_punctuation:
            chinese_count = sum(1 for c in para if '\u4e00' <= c <= '\u9fff')
            if chinese_count > 0:
                has_common_words = any(word in para for word in ['的', '是', '在', '了', '和', '与', '或', '中', '上', '下', '这', '那', '有', '为', '以', '及', '等', '也', '都', '就', '被', '把', '让', '给', '到', '从', '向', '对', '跟', '和', '同', '与', '比', '被', '把', '让', '给', '到', '从', '向', '对', '跟'])
                
                if has_common_words:
                    logger.debug(f"[_is_likely_heading] 包含常见连接词，不识别为标题: {para[:30]}...")
                    return False, None
                
                has_next_content = next_para and len(next_para.strip()) > 50
                has_prev_content = prev_para and len(prev_para.strip()) > 0
                
                if has_next_content or (not has_prev_content and not next_para):
                    logger.debug(f"[_is_likely_heading] 识别为中文标题: {para[:30]}...")
                    if para_len <= 15:
                        return True, 'h2'
                    elif para_len <= 30:
                        return True, 'h3'
                    elif para_len <= 50:
                        return True, 'h4'
                    else:
                        return True, 'h5'
            
            if prev_para and len(prev_para.strip()) > 0:
                if next_para and len(next_para.strip()) > 100:
                    logger.debug(f"[_is_likely_heading] 上下文判断为标题: {para[:30]}...")
                    return True, 'h3'
    
    logger.debug(f"[_is_likely_heading] 不识别为标题: {para[:30]}...")
    return False, None


def _is_list_item(para: str) -> tuple:
    if not para or len(para.strip()) == 0:
        return False, None, None
    
    if para.startswith('|') and para.endswith('|'):
        return False, None, None
    
    para = para.strip()
    
    ul_patterns = [
        r'^[•●○■□◆◇★☆►▶▸▹▻▪▫▬▭▮▯]\s*',
        r'^[-–—]\s+',
        r'^[*＊]\s+',
        r'^[√√✓✔✕✖✗✘]\s*',
        r'^\s*[•●○■□◆◇★☆►▶▸▹▻▪▫▬▭▮▯]\s*',
    ]
    
    for pattern in ul_patterns:
        match = re.match(pattern, para)
        if match:
            content = para[match.end():].strip()
            logger.debug(f"[_is_list_item] 识别为无序列表项: {content[:30]}...")
            return True, 'ul', content
    
    ol_patterns = [
        (r'^(\d+)[\.．、)\s]+', 'ol'),
        (r'^[（(](\d+)[)）]\s*', 'ol'),
        (r'^[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]\s*', 'ol'),
        (r'^[一二三四五六七八九十]+[、.．)\s]+', 'ol'),
        (r'^[（(][一二三四五六七八九十]+[)）]\s*', 'ol'),
    ]
    
    for pattern, list_type in ol_patterns:
        match = re.match(pattern, para)
        if match:
            content = para[match.end():].strip()
            logger.debug(f"[_is_list_item] 识别为有序列表项: {content[:30]}...")
            return True, list_type, content
    
    return False, None, None


def _is_table_row(para: str) -> bool:
    if not para or len(para.strip()) == 0:
        return False
    
    para = para.strip()
    
    if para.startswith('|') and para.endswith('|'):
        pipe_count = para.count('|')
        if pipe_count >= 3:
            logger.debug(f"[_is_table_row] 识别为表格行: {para[:50]}...")
            return True
    
    return False


def _convert_text_to_html(text: str) -> str:
    logger.info(f"[_convert_text_to_html] 开始转换，文本长度: {len(text)}")
    
    if not text or len(text.strip()) == 0:
        logger.warning("[_convert_text_to_html] 文本为空")
        return ""
    
    paragraphs = _smart_split_paragraphs(text)
    
    if not paragraphs:
        return ""
    
    html_parts = []
    in_list = False
    list_type = None
    in_table = False
    table_rows = []
    
    for i, para in enumerate(paragraphs):
        para = para.strip()
        if not para:
            continue
        
        prev_para = paragraphs[i-1].strip() if i > 0 else None
        next_para = paragraphs[i+1].strip() if i < len(paragraphs)-1 else None
        
        is_table_row = _is_table_row(para)
        
        if is_table_row:
            if not in_table:
                if in_list:
                    html_parts.append(f'</{list_type}>')
                    in_list = False
                    list_type = None
                in_table = True
                table_rows = []
            
            table_rows.append(para)
            continue
        
        if in_table:
            html_table = _convert_table_rows_to_html(table_rows)
            html_parts.append(html_table)
            in_table = False
            table_rows = []
        
        is_list, current_list_type, list_content = _is_list_item(para)
        
        if is_list:
            if not in_list or list_type != current_list_type:
                if in_list:
                    html_parts.append(f'</{list_type}>')
                list_type = current_list_type
                html_parts.append(f'<{list_type}>')
                in_list = True
            
            html_parts.append(f'<li>{list_content}</li>')
            continue
        
        if in_list:
            html_parts.append(f'</{list_type}>')
            in_list = False
            list_type = None
        
        is_heading, heading_level = _is_likely_heading(para, prev_para, next_para)
        
        if is_heading and heading_level:
            html_parts.append(f'<{heading_level}>{para}</{heading_level}>')
            continue
        
        para_lines = para.split('\n')
        if len(para_lines) > 1:
            para = '<br>'.join(para_lines)
        
        html_parts.append(f'<p>{para}</p>')
    
    if in_list:
        html_parts.append(f'</{list_type}>')
    
    if in_table and table_rows:
        html_table = _convert_table_rows_to_html(table_rows)
        html_parts.append(html_table)
    
    result = '\n'.join(html_parts)
    logger.info(f"[_convert_text_to_html] 转换完成，HTML长度: {len(result)}")
    logger.debug(f"[_convert_text_to_html] HTML内容: {result[:300]}...")
    
    return result


def _convert_table_rows_to_html(table_rows: list) -> str:
    if not table_rows:
        return ""
    
    logger.info(f"[_convert_table_rows_to_html] 转换 {len(table_rows)} 行表格")
    
    html_parts = ["<table border='1' cellpadding='5' cellspacing='0' style='border-collapse: collapse;'>"]
    
    for row_idx, row in enumerate(table_rows):
        cells = [cell.strip() for cell in row.split('|') if cell.strip()]
        
        if row_idx == 0:
            cell_tag = 'th'
        else:
            cell_tag = 'td'
        
        html_row = f"<tr>"
        for cell in cells:
            html_row += f"<{cell_tag}>{cell}</{cell_tag}>"
        html_row += "</tr>"
        html_parts.append(html_row)
    
    html_parts.append("</table>")
    
    result = '\n'.join(html_parts)
    logger.debug(f"[_convert_table_rows_to_html] 转换结果: {result[:200]}...")
    
    return result


def extract_simple_text(data: bytes) -> str:
    text_parts = []
    current_text = []
    
    for i in range(0, len(data) - 1, 2):
        try:
            char = data[i:i+2].decode('utf-16-le', errors='ignore')
            if char and (char.isprintable() or char in '\n\r\t'):
                current_text.append(char)
            elif current_text:
                text = ''.join(current_text)
                if len(text) > 3:
                    text_parts.append(text)
                current_text = []
        except:
            if current_text:
                text = ''.join(current_text)
                if len(text) > 3:
                    text_parts.append(text)
                current_text = []
    
    if current_text:
        text = ''.join(current_text)
        if len(text) > 3:
            text_parts.append(text)
    
    result = '\n'.join(text_parts)
    return _smart_clean_text(result)


def _try_decode_with_encoding(data: bytes, encoding: str) -> str:
    try:
        text = data.decode(encoding, errors='strict')
        return text
    except:
        try:
            text = data.decode(encoding, errors='ignore')
            return text
        except:
            return ""


def extract_text_from_doc(file_path: str) -> str:
    logger.warning(f"[extract_text_from_doc] 使用自定义解析器（备选方案）: {file_path}")
    
    if not olefile.isOleFile(file_path):
        logger.error("[extract_text_from_doc] 不是有效的 OLE 文件")
        return ""
    
    try:
        ole = olefile.OleFileIO(file_path)
        
        if not ole.exists('WordDocument'):
            logger.error("[extract_text_from_doc] 找不到 WordDocument 流")
            ole.close()
            return ""
        
        word_stream = ole.openstream('WordDocument')
        word_data = word_stream.read()
        word_stream.close()
        
        ole.close()
        
        extracted_texts = []
        
        try:
            fc_min = struct.unpack('<I', word_data[0x18:0x1C])[0]
            fc_max = struct.unpack('<I', word_data[0x1C:0x20])[0]
            
            logger.debug(f"[extract_text_from_doc] fc_min={fc_min}, fc_max={fc_max}")
            
            if 0 < fc_min < fc_max <= len(word_data):
                text_content = word_data[fc_min:fc_max]
                
                detected_encoding = _detect_encoding(text_content)
                
                encodings_to_try = [
                    detected_encoding,
                    'utf-16-le',
                    'utf-16-be',
                    'gbk',
                    'gb2312',
                    'gb18030',
                    'big5',
                    'shift_jis',
                    'euc-jp',
                    'euc-kr',
                    'utf-8',
                    'latin-1',
                    'cp1252',
                ]
                
                tried_encodings = set()
                for encoding in encodings_to_try:
                    if encoding and encoding.lower() not in tried_encodings:
                        tried_encodings.add(encoding.lower())
                        try:
                            text = _try_decode_with_encoding(text_content, encoding)
                            cleaned_text = _smart_clean_text(text)
                            if cleaned_text and len(cleaned_text.strip()) > 20:
                                if _is_highly_likely_valid_text(cleaned_text):
                                    logger.info(f"[extract_text_from_doc] 使用编码 {encoding} 提取成功，长度: {len(cleaned_text)}")
                                    return cleaned_text
                                elif _is_likely_valid_text(cleaned_text):
                                    extracted_texts.append(cleaned_text)
                        except:
                            continue
        except Exception as e:
            logger.warning(f"[extract_text_from_doc] 尝试提取时出错: {e}")
        
        for encoding in ['utf-16-le', 'utf-16-be', 'gbk', 'gb2312', 'gb18030', 'big5', 'utf-8', 'latin-1', 'cp1252']:
            try:
                text = _try_decode_with_encoding(word_data, encoding)
                cleaned_text = _smart_clean_text(text)
                if cleaned_text and len(cleaned_text.strip()) > 20:
                    if _is_highly_likely_valid_text(cleaned_text):
                        logger.info(f"[extract_text_from_doc] 使用编码 {encoding} 提取成功")
                        return cleaned_text
                    elif _is_likely_valid_text(cleaned_text):
                        extracted_texts.append(cleaned_text)
            except:
                continue
        
        simple_text = extract_simple_text(word_data)
        if simple_text and len(simple_text.strip()) > 20:
            if _is_highly_likely_valid_text(simple_text):
                return simple_text
            elif _is_likely_valid_text(simple_text):
                extracted_texts.append(simple_text)
        
        if extracted_texts:
            extracted_texts.sort(key=lambda x: (
                _is_highly_likely_valid_text(x),
                sum(1 for c in x if '\u4e00' <= c <= '\u9fff' or c.isalpha()),
                len(x)
            ), reverse=True)
            
            for text in extracted_texts:
                if _is_highly_likely_valid_text(text):
                    return text
            
            return extracted_texts[0]
        
        return ""
        
    except Exception as e:
        logger.error(f"[extract_text_from_doc] 提取文本时出错: {e}")
        import traceback
        traceback.print_exc()
        return ""


def parse_old_doc(file_path: str) -> tuple:
    logger.info(f"[parse_old_doc] 开始解析 .doc 文件: {file_path}")
    
    text = ""
    
    if HAS_ANTIWORD:
        logger.info("[parse_old_doc] 尝试使用 pyantiword")
        try:
            text = antiword_extract_text(file_path)
            logger.info(f"[parse_old_doc] pyantiword 提取成功，原始长度: {len(text)}")
            
            if text:
                text = text.strip()
                if len(text) > 10:
                    text = _smart_clean_text(text)
                    logger.info(f"[parse_old_doc] 清理后长度: {len(text)}")
                    
                    if _is_highly_likely_valid_text(text):
                        logger.info("[parse_old_doc] 文本有效，转换为HTML")
                        html_content = _convert_text_to_html(text)
                        return text, html_content
                    elif _is_likely_valid_text(text):
                        logger.info("[parse_old_doc] 文本可能有效，继续尝试其他方法")
                        pass
        except Exception as e:
            logger.error(f"[parse_old_doc] pyantiword 解析出错: {e}")
            import traceback
            traceback.print_exc()
    else:
        logger.warning("[parse_old_doc] pyantiword 不可用")
    
    logger.warning("[parse_old_doc] 回退到自定义解析器")
    text = extract_text_from_doc(file_path)
    
    if not text or len(text.strip()) < 10:
        logger.warning("[parse_old_doc] 自定义解析器提取失败，尝试简单提取")
        with open(file_path, 'rb') as f:
            data = f.read()
            text = extract_simple_text(data)
    
    if not text or not _is_likely_valid_text(text):
        logger.error("[parse_old_doc] 无法提取有效内容")
        return "[旧版Word文档 (.doc) - 无法提取有效内容]", "<p><em>旧版Word文档 (.doc) - 无法提取有效内容</em></p>"
    
    html_content = _convert_text_to_html(text)
    
    logger.info(f"[parse_old_doc] 解析完成，content长度: {len(text)}, html长度: {len(html_content)}")
    
    return text, html_content


def parse_docx_competitive(file_path: str) -> tuple:
    logger.info("=" * 60)
    logger.info(f"[parse_docx_competitive] 开始多解析器竞争解析: {file_path}")
    logger.info("=" * 60)
    
    competitor = get_multi_parser_competitor()
    best_result, all_results = competitor.parse_and_select(file_path)
    
    if best_result:
        logger.info(f"[parse_docx_competitive] 选择最优解析器: {best_result.parser_name}")
        logger.info(f"  - 得分: {best_result.score:.4f}")
        logger.info(f"  - 文本长度: {best_result.metrics.total_text_length}")
        logger.info(f"  - 标题数: {best_result.metrics.heading_count}")
        logger.info(f"  - 列表项: {best_result.metrics.list_item_count}")
        logger.info(f"  - 表格数: {best_result.metrics.table_count}")
        logger.info(f"  - 样式属性: {best_result.metrics.style_attributes_count}")
        
        return best_result.text_content, best_result.html_content
    else:
        logger.error("[parse_docx_competitive] 所有解析器都失败")
        return "[文档解析错误 - 所有解析器失败]", "<p><em>文档解析错误 - 所有解析器失败</em></p>"


def parse_docx_with_python_docx(file_path: str) -> tuple:
    parser = PythonDocxParser()
    result = parser.parse(file_path)
    
    if result.success:
        return result.text_content, result.html_content
    else:
        raise Exception(f"python-docx 解析失败: {result.error_message}")


def parse_docx_with_mammoth(file_path: str) -> tuple:
    parser = MammothParser()
    result = parser.parse(file_path)
    
    if result.success:
        return result.text_content, result.html_content
    else:
        raise Exception(f"mammoth 解析失败: {result.error_message}")


def parse_docx(file_path: str) -> tuple:
    logger.info(f"[parse_docx] 开始解析 .docx 文件: {file_path}")
    
    try:
        return parse_docx_competitive(file_path)
    except Exception as e:
        logger.error(f"[parse_docx] 多解析器竞争解析失败: {e}")
        
        logger.warning("[parse_docx] 回退到原始单解析器策略")
        try:
            return parse_docx_with_python_docx(file_path)
        except Exception as e2:
            logger.error(f"[parse_docx] python-docx 解析失败: {e2}")
            
            if HAS_MAMMOTH:
                logger.warning("[parse_docx] 回退到 mammoth 解析")
                try:
                    return parse_docx_with_mammoth(file_path)
                except Exception as e3:
                    logger.error(f"[parse_docx] mammoth 解析也失败: {e3}")
        
        raise Exception("所有解析方案都失败")


def parse_doc(file_path: str) -> tuple:
    logger.info("=" * 60)
    logger.info(f"[parse_doc] 开始解析文件: {file_path}")
    logger.info("=" * 60)
    
    file_path = Path(file_path)
    
    try:
        if is_docx_file(str(file_path)):
            logger.info("[parse_doc] 检测为 .docx 格式")
            return parse_docx(str(file_path))
        
        elif is_doc_file(str(file_path)):
            logger.info("[parse_doc] 检测为 .doc 格式")
            text, html = parse_old_doc(str(file_path))
            if text and _is_likely_valid_text(text):
                return text, html
            else:
                logger.warning("[parse_doc] .doc 解析结果无效")
                return f"[旧版Word文档 (.doc) - 无法提取有效内容]", "<p><em>旧版Word文档 (.doc) - 无法提取有效内容</em></p>"
        
        else:
            logger.warning("[parse_doc] 格式不明确，尝试各种解析器")
            try:
                return parse_docx(str(file_path))
            except BadZipFile:
                try:
                    if olefile.isOleFile(str(file_path)):
                        text, html = parse_old_doc(str(file_path))
                        if text and _is_likely_valid_text(text):
                            return text, html
                except:
                    pass
                
                logger.error("[parse_doc] 无法识别文件格式")
                return f"[不支持的Word文档格式]", "<p><em>不支持的Word文档格式</em></p>"
    except Exception as e:
        logger.error(f"[parse_doc] 解析出错: {e}")
        import traceback
        traceback.print_exc()
        return f"[文档解析错误: {str(e)}]", f"<p><em>文档解析错误: {str(e)}</em></p>"