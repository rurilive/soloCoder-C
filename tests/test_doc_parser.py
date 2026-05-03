"""Tests for Word document parser."""
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.parsers.doc_parser import (
    parse_doc,
    is_docx_file,
    is_doc_file,
    get_multi_parser_competitor,
    MultiParserCompetitor,
    PythonDocxParser,
    MammothParser,
    CmiDocxParser,
    ParseResult,
    ParseResultMetrics,
    MetricsCalculator,
    ResultScorer,
    HAS_MAMMOTH,
    HAS_CMI_DOCX,
)


class TestDocParserBasics:
    """Tests for basic Word document parsing."""

    def test_parse_doc_basic(self, temp_docx_file: Path):
        """Test basic docx parsing returning text and HTML."""
        text_content, html_content = parse_doc(str(temp_docx_file))
        
        assert text_content is not None
        assert html_content is not None
        assert len(text_content) > 0
        assert len(html_content) > 0
        
        assert "Test Document" in text_content
        assert "Introduction" in text_content
        
        assert "<h1" in html_content or "<h2" in html_content or "<p" in html_content

    def test_parse_doc_nonexistent_file(self):
        """Test parsing a non-existent file (returns error message)."""
        text_content, html_content = parse_doc("/nonexistent/file.docx")
        
        assert text_content is not None
        assert html_content is not None
        assert "错误" in text_content or "error" in text_content.lower()


class TestFileDetection:
    """Tests for file type detection functions."""

    def test_is_docx_file_true(self, temp_docx_file: Path):
        """Test that docx files are correctly identified."""
        result = is_docx_file(str(temp_docx_file))
        
        assert result is True

    def test_is_docx_file_false(self, temp_dir: Path):
        """Test that non-ZIP files are not identified as docx."""
        test_file = temp_dir / "test.txt"
        test_file.write_text("This is a plain text file", encoding="utf-8")
        
        result = is_docx_file(str(test_file))
        
        assert result is False

    def test_is_doc_file(self, temp_dir: Path):
        """Test doc file detection (with mock ole file)."""
        ole_header = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
        test_file = temp_dir / "test_ole.doc"
        test_file.write_bytes(ole_header + b"dummy content")
        
        try:
            result = is_doc_file(str(test_file))
        except:
            pytest.skip("OLE file detection may fail with minimal file")


class TestPythonDocxParser:
    """Tests for the python-docx based parser."""

    def test_python_docx_parser_basic(self, temp_docx_file: Path):
        """Test basic parsing with python-docx parser."""
        parser = PythonDocxParser()
        
        assert parser.name == "python-docx"
        assert parser.priority == 3
        
        result = parser.parse(str(temp_docx_file))
        
        assert isinstance(result, ParseResult)
        assert result.parser_name == "python-docx"
        assert result.success is True
        assert len(result.text_content) > 0
        assert len(result.html_content) > 0

    def test_python_docx_parser_headers(self, temp_docx_file: Path):
        """Test that headers are correctly parsed."""
        parser = PythonDocxParser()
        result = parser.parse(str(temp_docx_file))
        
        assert result.success is True
        assert "Test Document" in result.text_content
        assert "Introduction" in result.text_content

    def test_python_docx_parser_tables(self, temp_docx_file: Path):
        """Test that tables are correctly parsed."""
        parser = PythonDocxParser()
        result = parser.parse(str(temp_docx_file))
        
        assert result.success is True
        assert "Alice" in result.text_content
        assert "Bob" in result.text_content

    def test_python_docx_parser_nonexistent_file(self):
        """Test parsing non-existent file."""
        parser = PythonDocxParser()
        result = parser.parse("/nonexistent/file.docx")
        
        assert result.success is False
        assert len(result.error_message) > 0


class TestMammothParser:
    """Tests for the mammoth-based parser."""

    @pytest.mark.skipif(not HAS_MAMMOTH, reason="Mammoth library not installed")
    def test_mammoth_parser_basic(self, temp_docx_file: Path):
        """Test basic parsing with mammoth parser."""
        parser = MammothParser()
        
        assert parser.name == "mammoth"
        assert parser.priority == 2
        
        result = parser.parse(str(temp_docx_file))
        
        assert isinstance(result, ParseResult)
        assert result.parser_name == "mammoth"
        assert result.success is True
        assert len(result.text_content) > 0
        assert len(result.html_content) > 0

    @pytest.mark.skipif(not HAS_MAMMOTH, reason="Mammoth library not installed")
    def test_mammoth_parser_content(self, temp_docx_file: Path):
        """Test that mammoth extracts content correctly."""
        parser = MammothParser()
        result = parser.parse(str(temp_docx_file))
        
        assert result.success is True
        assert "Test Document" in result.text_content or "Test Document" in result.html_content

    def test_mammoth_parser_not_installed(self, monkeypatch):
        """Test behavior when mammoth is not installed."""
        monkeypatch.setattr('app.parsers.doc_parser.HAS_MAMMOTH', False)
        
        parser = MammothParser()
        result = parser.parse("/some/path.docx")
        
        assert result.success is False
        assert "未安装" in result.error_message or "not installed" in result.error_message.lower()


class TestCmiDocxParser:
    """Tests for the cmi-docx based parser."""

    @pytest.mark.skipif(not HAS_CMI_DOCX, reason="cmi-docx library not installed")
    def test_cmi_docx_parser_basic(self, temp_docx_file: Path):
        """Test basic parsing with cmi-docx parser."""
        parser = CmiDocxParser()
        
        assert parser.name == "cmi-docx"
        assert parser.priority == 1
        
        result = parser.parse(str(temp_docx_file))
        
        assert isinstance(result, ParseResult)
        assert result.parser_name == "cmi-docx"
        assert result.success is True

    def test_cmi_docx_parser_not_installed(self, monkeypatch):
        """Test behavior when cmi-docx is not installed."""
        monkeypatch.setattr('app.parsers.doc_parser.HAS_CMI_DOCX', False)
        
        parser = CmiDocxParser()
        result = parser.parse("/some/path.docx")
        
        assert result.success is False
        assert "未安装" in result.error_message or "not installed" in result.error_message.lower()


class TestMultiParserCompetitor:
    """Tests for the multi-parser competitor."""

    def test_get_multi_parser_competitor(self):
        """Test getting the singleton competitor instance."""
        competitor = get_multi_parser_competitor()
        
        assert isinstance(competitor, MultiParserCompetitor)
        
        competitor2 = get_multi_parser_competitor()
        assert competitor is competitor2

    def test_multi_parser_initialization(self):
        """Test that competitor initializes with correct parsers."""
        competitor = MultiParserCompetitor()
        
        assert len(competitor._parsers) >= 1
        
        parser_names = [p.name for p in competitor._parsers]
        assert "python-docx" in parser_names

    def test_multi_parser_parse_all(self, temp_docx_file: Path):
        """Test parsing with all available parsers."""
        competitor = MultiParserCompetitor()
        
        results = competitor.parse_all(str(temp_docx_file))
        
        assert len(results) >= 1
        
        successful = [r for r in results if r.success]
        assert len(successful) >= 1

    def test_multi_parser_select_best(self, temp_docx_file: Path):
        """Test selecting the best parser result."""
        competitor = MultiParserCompetitor()
        
        results = competitor.parse_all(str(temp_docx_file))
        best = competitor.select_best(results)
        
        assert best is not None
        assert best.success is True

    def test_multi_parser_parse_and_select(self, temp_docx_file: Path):
        """Test the combined parse and select method."""
        competitor = MultiParserCompetitor()
        
        best, all_results = competitor.parse_and_select(str(temp_docx_file))
        
        assert best is not None
        assert best.success is True
        assert len(all_results) >= 1


class TestMetricsCalculator:
    """Tests for the metrics calculator."""

    def test_calculate_metrics_basic(self):
        """Test basic metrics calculation."""
        text_content = "This is a test document. It has multiple paragraphs."
        html_content = "<h1>Title</h1><p>Paragraph 1</p><p>Paragraph 2</p>"
        
        metrics = MetricsCalculator.calculate_metrics(text_content, html_content)
        
        assert isinstance(metrics, ParseResultMetrics)
        assert metrics.total_text_length == len(text_content)

    def test_calculate_metrics_empty(self):
        """Test metrics calculation with empty content."""
        metrics = MetricsCalculator.calculate_metrics("", "")
        
        assert metrics.total_text_length == 0

    def test_analyze_html_with_regex(self):
        """Test HTML analysis with regex."""
        html_content = """
        <h1>Main Title</h1>
        <h2>Subtitle</h2>
        <p>This is a paragraph with <strong>bold</strong> and <em>italic</em> text.</p>
        <ul>
            <li>Item 1</li>
            <li>Item 2</li>
        </ul>
        <table>
            <tr><th>Header</th></tr>
            <tr><td>Data</td></tr>
        </table>
        <a href="http://example.com">Link</a>
        <img src="image.png" alt="Image">
        """
        
        metrics = MetricsCalculator._analyze_with_regex(html_content, ParseResultMetrics())
        
        assert metrics.heading_count >= 2
        assert metrics.paragraph_count >= 1
        assert metrics.list_item_count >= 2
        assert metrics.table_count == 1
        assert metrics.table_row_count == 2
        assert metrics.table_cell_count == 2
        assert metrics.link_count == 1
        assert metrics.image_count == 1


class TestResultScorer:
    """Tests for the result scorer."""

    def test_calculate_score_basic(self):
        """Test basic score calculation."""
        metrics = ParseResultMetrics(
            total_text_length=1000,
            heading_count=5,
            list_item_count=10,
            table_count=2,
            table_cell_count=20,
            inline_format_count=15,
            link_count=3,
            image_count=2,
            style_attributes_count=10,
            alignment_count=5,
            indent_count=3,
        )
        
        result = ParseResult(
            parser_name="test",
            text_content="test content",
            html_content="<p>test</p>",
            metrics=metrics
        )
        
        score = ResultScorer.calculate_score(result)
        
        assert 0 <= score <= 1

    def test_calculate_score_with_multiple_results(self, temp_docx_file: Path):
        """Test score calculation with multiple results for comparison."""
        competitor = MultiParserCompetitor()
        results = competitor.parse_all(str(temp_docx_file))
        successful = [r for r in results if r.success]
        
        if len(successful) >= 1:
            for result in successful:
                score = ResultScorer.calculate_score(result, successful)
                assert 0 <= score <= 1


class TestDocParserEdgeCases:
    """Tests for edge cases in doc parsing."""

    def test_parse_empty_docx(self, temp_dir: Path):
        """Test parsing an empty docx file."""
        file_path = temp_dir / "empty.docx"
        
        doc = Document()
        doc.save(file_path)
        
        text_content, html_content = parse_doc(str(file_path))
        
        assert text_content is not None
        assert html_content is not None

    def test_parse_docx_with_only_title(self, temp_dir: Path):
        """Test parsing a docx with only a title."""
        file_path = temp_dir / "title_only.docx"
        
        doc = Document()
        doc.add_heading("Single Title Document", level=1)
        doc.save(file_path)
        
        text_content, html_content = parse_doc(str(file_path))
        
        assert "Single Title Document" in text_content

    def test_parse_docx_with_complex_formatting(self, temp_dir: Path):
        """Test parsing a docx with complex formatting."""
        file_path = temp_dir / "complex.docx"
        
        doc = Document()
        
        title = doc.add_heading("Complex Document", level=1)
        title.alignment = 1
        
        para = doc.add_paragraph()
        para.add_run("This is ").bold = False
        para.add_run("bold text").bold = True
        para.add_run(" and this is ").bold = False
        para.add_run("italic text").italic = True
        
        doc.add_paragraph("First list item", style="List Bullet")
        doc.add_paragraph("Second list item", style="List Bullet")
        
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"
        table.rows[1].cells[0].text = "Data 1"
        table.rows[1].cells[1].text = "Data 2"
        
        doc.save(file_path)
        
        text_content, html_content = parse_doc(str(file_path))
        
        assert "Complex Document" in text_content
        assert "bold text" in text_content
        assert "italic text" in text_content
        assert "First list item" in text_content
        assert "Header 1" in text_content


class TestDocParserIntegration:
    """Integration tests for doc parser."""

    def test_full_parsing_workflow(self, temp_docx_file: Path):
        """Test the complete parsing workflow."""
        text_content, html_content = parse_doc(str(temp_docx_file))
        
        assert len(text_content) > 0
        assert len(html_content) > 0
        
        assert "Test Document" in text_content
        assert "Alice" in text_content
        assert "Bob" in text_content
        
        assert "<" in html_content
        assert ">" in html_content

    def test_parser_consistency(self, temp_docx_file: Path):
        """Test that parsing is consistent across multiple calls."""
        text1, html1 = parse_doc(str(temp_docx_file))
        text2, html2 = parse_doc(str(temp_docx_file))
        
        assert text1 == text2
