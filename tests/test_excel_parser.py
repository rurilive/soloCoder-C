"""Tests for Excel parser."""
import sys
from pathlib import Path
from datetime import datetime

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.parsers.excel_parser import (
    parse_excel,
    parse_excel_file,
    update_cell_in_excel,
    update_multiple_cells,
    add_new_sheet,
    delete_sheet,
    get_sheet_metadata,
    get_sheet_data_paginated,
    get_sheet_data_as_json,
    ParseResult,
    ExcelSheet,
    ExcelCell,
)


class TestExcelParserBasics:
    """Tests for basic Excel parsing functionality."""

    def test_parse_excel_basic(self, temp_excel_file: Path):
        """Test basic excel parsing returning text and HTML."""
        text_content, html_content = parse_excel(str(temp_excel_file))
        
        assert text_content is not None
        assert html_content is not None
        assert len(text_content) > 0
        assert len(html_content) > 0
        
        assert "Sheet1" in text_content or "Sheet2" in text_content
        assert "Alice" in text_content
        assert "Bob" in text_content
        
        assert "<table" in html_content
        assert "<th" in html_content or "<td" in html_content

    def test_parse_excel_file_detailed(self, temp_excel_file: Path):
        """Test detailed excel parsing with full structure."""
        result = parse_excel_file(str(temp_excel_file))
        
        assert isinstance(result, ParseResult)
        assert result.success is True
        assert len(result.sheets) == 2
        
        sheet1 = result.sheets[0]
        assert isinstance(sheet1, ExcelSheet)
        assert sheet1.name == "Sheet1"
        assert sheet1.max_row == 4
        assert sheet1.max_col == 4
        
        assert len(sheet1.rows) == 4
        assert len(sheet1.rows[0]) == 4
        
        assert sheet1.rows[0][0].value == "Name"
        assert sheet1.rows[1][0].value == "Alice"
        assert sheet1.rows[2][0].value == "Bob"
        
        sheet2 = result.sheets[1]
        assert sheet2.name == "Sheet2"
        assert sheet2.max_row == 3
        assert sheet2.max_col == 3

    def test_parse_excel_file_nonexistent(self):
        """Test parsing a non-existent file."""
        result = parse_excel_file("/nonexistent/file.xlsx")
        
        assert result.success is False
        assert "无法加载" in result.error_message or "error" in result.error_message.lower()


class TestExcelMetadata:
    """Tests for Excel metadata extraction."""

    def test_get_sheet_metadata(self, temp_excel_file: Path):
        """Test getting sheet metadata without loading full data."""
        metadata = get_sheet_metadata(str(temp_excel_file))
        
        assert metadata["success"] is True
        assert metadata["sheet_count"] == 2
        
        sheets = metadata["sheets"]
        assert len(sheets) == 2
        
        sheet1 = sheets[0]
        assert sheet1["name"] == "Sheet1"
        assert sheet1["max_row"] == 4
        assert sheet1["max_col"] == 4
        
        sheet2 = sheets[1]
        assert sheet2["name"] == "Sheet2"
        assert sheet2["max_row"] == 3
        assert sheet2["max_col"] == 3

    def test_get_sheet_metadata_nonexistent(self):
        """Test getting metadata for non-existent file."""
        metadata = get_sheet_metadata("/nonexistent/file.xlsx")
        
        assert metadata["success"] is False
        assert "error" in metadata


class TestExcelDataExtraction:
    """Tests for Excel data extraction functionality."""

    def test_get_sheet_data_as_json_all_sheets(self, temp_excel_file: Path):
        """Test getting all sheets data as JSON."""
        data = get_sheet_data_as_json(str(temp_excel_file))
        
        assert data["success"] is True
        assert data["sheet_count"] == 2
        assert data["is_large_file"] is False
        
        sheets = data["sheets"]
        assert len(sheets) == 2
        
        sheet1 = sheets[0]
        assert sheet1["name"] == "Sheet1"
        assert sheet1["max_row"] == 4
        assert sheet1["max_col"] == 4
        assert sheet1["is_paginated"] is False
        assert len(sheet1["data"]) == 4
        
        assert sheet1["data"][0][0]["value"] == "Name"
        assert sheet1["data"][1][0]["value"] == "Alice"
        assert sheet1["data"][2][0]["value"] == "Bob"

    def test_get_sheet_data_as_json_single_sheet(self, temp_excel_file: Path):
        """Test getting data for a specific sheet."""
        data = get_sheet_data_as_json(str(temp_excel_file), sheet_name="Sheet2")
        
        assert data["success"] is True
        assert data["sheet_count"] == 1
        
        sheets = data["sheets"]
        assert len(sheets) == 1
        assert sheets[0]["name"] == "Sheet2"
        assert sheets[0]["max_row"] == 3
        assert sheets[0]["max_col"] == 3

    def test_get_sheet_data_as_json_nonexistent_sheet(self, temp_excel_file: Path):
        """Test getting data for a non-existent sheet."""
        data = get_sheet_data_as_json(str(temp_excel_file), sheet_name="NonexistentSheet")
        
        assert data["success"] is True
        assert data["sheet_count"] == 0
        assert len(data["sheets"]) == 0


class TestExcelPagination:
    """Tests for Excel pagination functionality."""

    def test_get_sheet_data_paginated_basic(self, temp_excel_file: Path):
        """Test basic pagination functionality."""
        data = get_sheet_data_paginated(
            str(temp_excel_file),
            sheet_name="Sheet1",
            start_row=1,
            end_row=2,
            include_styles=False
        )
        
        assert data["success"] is True
        assert data["sheet_name"] == "Sheet1"
        assert data["total_rows"] == 4
        assert data["total_cols"] == 4
        assert data["start_row"] == 1
        assert data["end_row"] == 2
        assert data["row_count"] == 2
        assert len(data["data"]) == 2

    def test_get_sheet_data_paginated_middle_range(self, temp_excel_file: Path):
        """Test pagination with middle range."""
        data = get_sheet_data_paginated(
            str(temp_excel_file),
            sheet_name="Sheet1",
            start_row=2,
            end_row=3,
            include_styles=False
        )
        
        assert data["success"] is True
        assert data["start_row"] == 2
        assert data["end_row"] == 3
        assert data["row_count"] == 2
        
        assert data["data"][0][0]["value"] == "Alice"
        assert data["data"][1][0]["value"] == "Bob"

    def test_get_sheet_data_paginated_out_of_range(self, temp_excel_file: Path):
        """Test pagination with out-of-range rows."""
        data = get_sheet_data_paginated(
            str(temp_excel_file),
            sheet_name="Sheet1",
            start_row=100,
            end_row=200,
            include_styles=False
        )
        
        assert data["success"] is True
        assert data["row_count"] == 0
        assert len(data["data"]) == 0

    def test_get_sheet_data_paginated_nonexistent_sheet(self, temp_excel_file: Path):
        """Test pagination for non-existent sheet."""
        data = get_sheet_data_paginated(
            str(temp_excel_file),
            sheet_name="NonexistentSheet",
            start_row=1,
            end_row=10,
            include_styles=False
        )
        
        assert data["success"] is False
        assert "不存在" in data["error"] or "exist" in data["error"].lower()

    def test_get_sheet_data_paginated_with_styles(self, temp_excel_file: Path):
        """Test pagination with style information."""
        data = get_sheet_data_paginated(
            str(temp_excel_file),
            sheet_name="Sheet1",
            start_row=1,
            end_row=1,
            include_styles=True
        )
        
        assert data["success"] is True
        assert len(data["data"]) == 1
        
        first_cell = data["data"][0][0]
        assert "style" in first_cell
        assert "font" in first_cell["style"]


class TestExcelCellUpdates:
    """Tests for Excel cell update functionality."""

    def test_update_cell_in_excel_basic(self, temp_excel_file: Path):
        """Test updating a single cell."""
        wb = load_workbook(temp_excel_file)
        original_value = wb["Sheet1"]["A2"].value
        wb.close()
        
        assert original_value == "Alice"
        
        success = update_cell_in_excel(
            str(temp_excel_file),
            "Sheet1",
            row=2,
            col=1,
            value="Alicia"
        )
        
        assert success is True
        
        wb = load_workbook(temp_excel_file)
        new_value = wb["Sheet1"]["A2"].value
        wb.close()
        
        assert new_value == "Alicia"

    def test_update_cell_in_excel_nonexistent_sheet(self, temp_excel_file: Path):
        """Test updating a cell in non-existent sheet."""
        success = update_cell_in_excel(
            str(temp_excel_file),
            "NonexistentSheet",
            row=1,
            col=1,
            value="Test"
        )
        
        assert success is False

    def test_update_multiple_cells_basic(self, temp_excel_file: Path):
        """Test updating multiple cells at once."""
        updates = [
            {"sheet_name": "Sheet1", "row": 2, "col": 1, "value": "Alicia"},
            {"sheet_name": "Sheet1", "row": 2, "col": 2, "value": 31},
            {"sheet_name": "Sheet2", "row": 2, "col": 2, "value": 1099},
        ]
        
        success = update_multiple_cells(str(temp_excel_file), updates)
        
        assert success is True
        
        wb = load_workbook(temp_excel_file)
        assert wb["Sheet1"]["A2"].value == "Alicia"
        assert wb["Sheet1"]["B2"].value == 31
        assert wb["Sheet2"]["B2"].value == 1099
        wb.close()

    def test_update_multiple_cells_mixed_sheets(self, temp_excel_file: Path):
        """Test updating cells with mixed valid and invalid sheets."""
        updates = [
            {"sheet_name": "Sheet1", "row": 2, "col": 1, "value": "Alicia"},
            {"sheet_name": "NonexistentSheet", "row": 1, "col": 1, "value": "Test"},
        ]
        
        success = update_multiple_cells(str(temp_excel_file), updates)
        
        assert success is True
        
        wb = load_workbook(temp_excel_file)
        assert wb["Sheet1"]["A2"].value == "Alicia"
        wb.close()


class TestExcelSheetManagement:
    """Tests for Excel sheet management functionality."""

    def test_add_new_sheet_basic(self, temp_excel_file: Path):
        """Test adding a new sheet."""
        wb = load_workbook(temp_excel_file)
        original_count = len(wb.sheetnames)
        wb.close()
        
        assert original_count == 2
        
        success = add_new_sheet(str(temp_excel_file), "NewSheet")
        
        assert success is True
        
        wb = load_workbook(temp_excel_file)
        assert len(wb.sheetnames) == 3
        assert "NewSheet" in wb.sheetnames
        wb.close()

    def test_add_new_sheet_duplicate_name(self, temp_excel_file: Path):
        """Test adding a sheet with duplicate name."""
        success = add_new_sheet(str(temp_excel_file), "Sheet1")
        
        assert success is False
        
        wb = load_workbook(temp_excel_file)
        assert len(wb.sheetnames) == 2
        wb.close()

    def test_delete_sheet_basic(self, temp_excel_file: Path):
        """Test deleting a sheet."""
        wb = load_workbook(temp_excel_file)
        original_count = len(wb.sheetnames)
        wb.close()
        
        assert original_count == 2
        
        success = delete_sheet(str(temp_excel_file), "Sheet2")
        
        assert success is True
        
        wb = load_workbook(temp_excel_file)
        assert len(wb.sheetnames) == 1
        assert "Sheet2" not in wb.sheetnames
        wb.close()

    def test_delete_sheet_nonexistent(self, temp_excel_file: Path):
        """Test deleting a non-existent sheet."""
        success = delete_sheet(str(temp_excel_file), "NonexistentSheet")
        
        assert success is False
        
        wb = load_workbook(temp_excel_file)
        assert len(wb.sheetnames) == 2
        wb.close()

    def test_delete_sheet_last_one(self, temp_excel_file: Path):
        """Test deleting the last sheet should fail."""
        wb = load_workbook(temp_excel_file)
        wb.remove(wb["Sheet2"])
        wb.save(temp_excel_file)
        wb.close()
        
        wb = load_workbook(temp_excel_file)
        assert len(wb.sheetnames) == 1
        wb.close()
        
        success = delete_sheet(str(temp_excel_file), "Sheet1")
        
        assert success is False
        
        wb = load_workbook(temp_excel_file)
        assert len(wb.sheetnames) == 1
        wb.close()


class TestExcelCellTypes:
    """Tests for Excel cell type detection."""

    def test_cell_type_detection(self, temp_excel_file: Path):
        """Test that cell types are correctly detected."""
        result = parse_excel_file(str(temp_excel_file))
        
        assert result.success is True
        
        sheet1 = result.sheets[0]
        
        header_cell = sheet1.rows[0][0]
        assert header_cell.data_type == "string"
        
        name_cell = sheet1.rows[1][0]
        assert name_cell.data_type == "string"
        
        age_cell = sheet1.rows[1][1]
        assert age_cell.data_type == "number"
        
        salary_cell = sheet1.rows[1][3]
        assert salary_cell.data_type == "number"


class TestExcelMergedCells:
    """Tests for Excel merged cells handling."""

    def test_merged_cells_detection(self, temp_dir: Path):
        """Test that merged cells are correctly detected."""
        from openpyxl import Workbook
        
        file_path = temp_dir / "merged_test.xlsx"
        
        wb = Workbook()
        ws = wb.active
        
        ws.merge_cells('A1:C1')
        ws['A1'] = "Merged Header"
        
        ws.append(["Col1", "Col2", "Col3"])
        ws.append(["Data1", "Data2", "Data3"])
        
        wb.save(file_path)
        
        result = parse_excel_file(str(file_path))
        
        assert result.success is True
        
        sheet = result.sheets[0]
        assert len(sheet.merged_cells) == 1
        
        merged_info = sheet.merged_cells[0]
        assert merged_info["min_row"] == 1
        assert merged_info["min_col"] == 1
        assert merged_info["max_row"] == 1
        assert merged_info["max_col"] == 3
        
        assert sheet.rows[0][0].is_merged is True
        assert sheet.rows[0][0].merge_master is True
        
        assert sheet.rows[0][1].is_merged is True
        assert sheet.rows[0][1].merge_master is False


class TestExcelEdgeCases:
    """Tests for edge cases in Excel parsing."""

    def test_empty_excel_file(self, temp_dir: Path):
        """Test parsing an empty Excel file."""
        from openpyxl import Workbook
        
        file_path = temp_dir / "empty.xlsx"
        
        wb = Workbook()
        ws = wb.active
        ws.title = "EmptySheet"
        wb.save(file_path)
        
        result = parse_excel_file(str(file_path))
        
        assert result.success is True
        assert len(result.sheets) == 1
        
        sheet = result.sheets[0]
        assert sheet.name == "EmptySheet"
        assert sheet.max_row == 1
        assert sheet.max_col == 1

    def test_excel_file_with_formulas(self, temp_dir: Path):
        """Test parsing Excel file with formulas."""
        from openpyxl import Workbook
        
        file_path = temp_dir / "formula_test.xlsx"
        
        wb = Workbook()
        ws = wb.active
        
        ws['A1'] = "Price"
        ws['B1'] = "Quantity"
        ws['C1'] = "Total"
        
        ws['A2'] = 100
        ws['B2'] = 5
        ws['C2'] = "=A2*B2"
        
        wb.save(file_path)
        
        result = parse_excel_file(str(file_path))
        
        assert result.success is True
        
        sheet = result.sheets[0]
        
        assert sheet.rows[0][0].value == "Price"
        assert sheet.rows[1][0].value == 100
        assert sheet.rows[1][1].value == 5
        
        formula_cell = sheet.rows[1][2]
        assert formula_cell.data_type == "formula" or formula_cell.data_type == "string"


class TestExcelHtmlGeneration:
    """Tests for Excel HTML generation."""

    def test_html_generation_basic(self, temp_excel_file: Path):
        """Test basic HTML generation from Excel."""
        result = parse_excel_file(str(temp_excel_file))
        
        assert result.success is True
        assert len(result.html_content) > 0
        
        assert "<h2>工作表: Sheet1</h2>" in result.html_content or "Sheet1" in result.html_content
        assert "<table" in result.html_content
        assert "<th" in result.html_content or "<td" in result.html_content
        assert "</table>" in result.html_content

    def test_html_contains_table_structure(self, temp_excel_file: Path):
        """Test that HTML contains proper table structure."""
        _, html_content = parse_excel(str(temp_excel_file))
        
        assert "<table" in html_content
        assert "<tr" in html_content
        assert "</tr>" in html_content
        assert "<td" in html_content or "<th" in html_content
        assert "</table>" in html_content

    def test_html_escaping(self, temp_dir: Path):
        """Test that special characters are properly escaped in HTML."""
        from openpyxl import Workbook
        
        file_path = temp_dir / "escape_test.xlsx"
        
        wb = Workbook()
        ws = wb.active
        
        ws['A1'] = 'Text with <script>alert("xss")</script>'
        ws['A2'] = 'Text with & and < and > characters'
        
        wb.save(file_path)
        
        _, html_content = parse_excel(str(file_path))
        
        assert '<script>' not in html_content
        assert '&lt;' in html_content or '&amp;' in html_content
