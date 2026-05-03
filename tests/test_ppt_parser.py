"""Tests for PowerPoint parser."""
import sys
import shutil
from pathlib import Path
from typing import Generator

import pytest
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.parsers.ppt_parser import (
    parse_ppt,
    parse_ppt_file,
    get_ppt_metadata,
    get_slide_data_paginated,
    get_ppt_data_as_json,
    update_text_in_slide,
    update_multiple_texts,
    add_new_slide,
    delete_slide,
    reorder_slides,
    ParseResult,
    PresentationInfo,
    SlideInfo,
)


class TestPptParserBasics:
    """Tests for basic PowerPoint parsing functionality."""

    def test_parse_ppt_basic(self, temp_pptx_file: Path):
        """Test basic pptx parsing returning text and HTML."""
        text_content, html_content = parse_ppt(str(temp_pptx_file))
        
        assert text_content is not None
        assert html_content is not None
        assert len(text_content) > 0
        assert len(html_content) > 0
        
        assert "Test Presentation" in text_content
        assert "Created for testing purposes" in text_content

    def test_parse_ppt_nonexistent_file(self):
        """Test parsing a non-existent file."""
        with pytest.raises(Exception):
            parse_ppt("/nonexistent/file.pptx")

    def test_parse_ppt_file_detailed(self, temp_pptx_file: Path):
        """Test detailed pptx parsing with full structure."""
        result = parse_ppt_file(str(temp_pptx_file))
        
        assert isinstance(result, ParseResult)
        assert result.success is True
        assert result.presentation is not None
        assert isinstance(result.presentation, PresentationInfo)
        assert result.presentation.slide_count == 3
        assert len(result.text_content) > 0
        assert len(result.html_content) > 0

    def test_parse_ppt_file_nonexistent(self):
        """Test parsing a non-existent file with detailed parser."""
        result = parse_ppt_file("/nonexistent/file.pptx")
        
        assert result.success is False
        assert "无法加载PPT文件" in result.error_message


class TestPptParserSlides:
    """Tests for slide parsing functionality."""

    def test_parse_ppt_slide_count(self, temp_pptx_file: Path):
        """Test that correct number of slides are parsed."""
        text_content, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "幻灯片 1" in text_content or "Slide 1" in text_content
        assert "幻灯片 2" in text_content or "Slide 2" in text_content
        assert "幻灯片 3" in text_content or "Slide 3" in text_content

    def test_parse_ppt_title_slide(self, temp_pptx_file: Path):
        """Test that title slide content is correctly parsed."""
        text_content, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "Test Presentation" in text_content
        assert "Created for testing purposes" in text_content

    def test_parse_ppt_bullet_slide(self, temp_pptx_file: Path):
        """Test that bullet points are correctly parsed."""
        text_content, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "First bullet point" in text_content
        assert "Second bullet point" in text_content
        assert "Third bullet point" in text_content

    def test_parse_ppt_table_slide(self, temp_pptx_file: Path):
        """Test that tables in slides are correctly parsed."""
        text_content, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "Name" in text_content
        assert "Age" in text_content
        assert "City" in text_content
        assert "Alice" in text_content
        assert "Bob" in text_content

    def test_slide_info_structure(self, temp_pptx_file: Path):
        """Test that slide info structure is correct."""
        result = parse_ppt_file(str(temp_pptx_file))
        
        assert len(result.presentation.slides) == 3
        
        for slide in result.presentation.slides:
            assert isinstance(slide, SlideInfo)
            assert slide.slide_idx >= 1
            assert len(slide.shapes) >= 0


class TestPptParserHtmlGeneration:
    """Tests for HTML generation from PowerPoint."""

    def test_ppt_html_contains_slide_headers(self, temp_pptx_file: Path):
        """Test that HTML contains slide headers."""
        _, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "<h2" in html_content
        assert "幻灯片 1" in html_content or "Slide 1" in html_content

    def test_ppt_html_contains_paragraphs(self, temp_pptx_file: Path):
        """Test that HTML contains paragraph tags."""
        _, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "<p" in html_content
        assert "</p>" in html_content

    def test_ppt_html_contains_tables(self, temp_pptx_file: Path):
        """Test that HTML contains table tags for table slides."""
        _, html_content = parse_ppt(str(temp_pptx_file))
        
        assert "<table" in html_content
        assert "<tr" in html_content
        assert "<td" in html_content
        assert "</table>" in html_content


class TestPptParserEdgeCases:
    """Tests for edge cases in PowerPoint parsing."""

    def test_parse_empty_pptx(self, temp_dir: Path):
        """Test parsing an empty PowerPoint file."""
        file_path = temp_dir / "empty.pptx"
        
        prs = Presentation()
        prs.save(file_path)
        
        text_content, html_content = parse_ppt(str(file_path))
        
        assert text_content is not None
        assert html_content is not None

    def test_parse_pptx_with_only_title(self, temp_dir: Path):
        """Test parsing a PowerPoint with only a title slide."""
        file_path = temp_dir / "title_only.pptx"
        
        prs = Presentation()
        
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Single Slide Presentation"
        subtitle.text = "Only one slide"
        
        prs.save(file_path)
        
        text_content, html_content = parse_ppt(str(file_path))
        
        assert "Single Slide Presentation" in text_content
        assert "Only one slide" in text_content

    def test_parse_pptx_with_multiple_shapes(self, temp_dir: Path):
        """Test parsing a PowerPoint with multiple text shapes."""
        file_path = temp_dir / "multiple_shapes.pptx"
        
        prs = Presentation()
        
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        title.text = "Main Title"
        
        bullet_slide_layout = prs.slide_layouts[1]
        slide2 = prs.slides.add_slide(bullet_slide_layout)
        
        title_shape = slide2.shapes.title
        title_shape.text = "Bullet Slide"
        
        body_shape = slide2.shapes.placeholders[1]
        tf = body_shape.text_frame
        tf.text = "First point"
        
        p = tf.add_paragraph()
        p.text = "Second point"
        p.level = 1
        
        p = tf.add_paragraph()
        p.text = "Third point"
        p.level = 2
        
        prs.save(file_path)
        
        text_content, html_content = parse_ppt(str(file_path))
        
        assert "Main Title" in text_content
        assert "Bullet Slide" in text_content
        assert "First point" in text_content
        assert "Second point" in text_content
        assert "Third point" in text_content


class TestPptParserIntegration:
    """Integration tests for PowerPoint parser."""

    def test_full_parsing_workflow(self, temp_pptx_file: Path):
        """Test the complete parsing workflow."""
        text_content, html_content = parse_ppt(str(temp_pptx_file))
        
        assert len(text_content) > 0
        assert len(html_content) > 0
        
        assert "Test Presentation" in text_content
        assert "Second Slide" in text_content
        assert "Table Slide" in text_content
        
        assert "<h2" in html_content
        assert "<p" in html_content
        assert "<table" in html_content

    def test_parser_consistency(self, temp_pptx_file: Path):
        """Test that parsing is consistent across multiple calls."""
        text1, html1 = parse_ppt(str(temp_pptx_file))
        text2, html2 = parse_ppt(str(temp_pptx_file))
        
        assert text1 == text2
        assert html1 == html2


class TestPptMetadata:
    """Tests for PowerPoint metadata extraction."""

    def test_get_ppt_metadata_basic(self, temp_pptx_file: Path):
        """Test basic metadata extraction."""
        metadata = get_ppt_metadata(str(temp_pptx_file))
        
        assert metadata['success'] is True
        assert metadata['slide_count'] == 3
        assert 'slide_width' in metadata
        assert 'slide_height' in metadata
        assert len(metadata['slides']) == 3

    def test_get_ppt_metadata_slide_info(self, temp_pptx_file: Path):
        """Test that metadata contains slide information."""
        metadata = get_ppt_metadata(str(temp_pptx_file))
        
        for slide in metadata['slides']:
            assert 'slide_idx' in slide
            assert 'shape_count' in slide
            assert 'layout_name' in slide

    def test_get_ppt_metadata_nonexistent(self):
        """Test metadata extraction for non-existent file."""
        metadata = get_ppt_metadata("/nonexistent/file.pptx")
        
        assert metadata['success'] is False
        assert 'error' in metadata


class TestPptPaginatedData:
    """Tests for paginated slide data extraction."""

    def test_get_slide_data_paginated_basic(self, temp_pptx_file: Path):
        """Test basic paginated data extraction."""
        data = get_slide_data_paginated(str(temp_pptx_file), start_slide=1, end_slide=2)
        
        assert data['success'] is True
        assert data['total_slides'] == 3
        assert data['start_slide'] == 1
        assert data['end_slide'] == 2
        assert data['slide_count'] == 2
        assert len(data['slides']) == 2

    def test_get_slide_data_paginated_single_slide(self, temp_pptx_file: Path):
        """Test extracting a single slide."""
        data = get_slide_data_paginated(str(temp_pptx_file), start_slide=2, end_slide=2)
        
        assert data['success'] is True
        assert data['slide_count'] == 1
        assert len(data['slides']) == 1
        assert data['slides'][0]['slide_idx'] == 2

    def test_get_slide_data_paginated_out_of_range(self, temp_pptx_file: Path):
        """Test extracting slides that are out of range."""
        data = get_slide_data_paginated(str(temp_pptx_file), start_slide=10, end_slide=20)
        
        assert data['success'] is True
        assert data['slide_count'] == 0
        assert len(data['slides']) == 0

    def test_get_slide_data_paginated_with_notes(self, temp_pptx_file: Path):
        """Test extracting data with notes enabled."""
        data = get_slide_data_paginated(
            str(temp_pptx_file), 
            start_slide=1, 
            end_slide=1,
            include_notes=True
        )
        
        assert data['success'] is True
        assert len(data['slides']) == 1

    def test_get_slide_data_paginated_nonexistent(self):
        """Test paginated data extraction for non-existent file."""
        data = get_slide_data_paginated("/nonexistent/file.pptx", start_slide=1, end_slide=1)
        
        assert data['success'] is False
        assert 'error' in data


class TestPptDataAsJson:
    """Tests for getting PowerPoint data as JSON."""

    def test_get_ppt_data_as_json_basic(self, temp_pptx_file: Path):
        """Test basic JSON data extraction."""
        data = get_ppt_data_as_json(str(temp_pptx_file))
        
        assert data['success'] is True
        assert data['slide_count'] == 3
        assert data['is_paginated'] is False
        assert len(data['slides']) == 3

    def test_get_ppt_data_as_json_with_range(self, temp_pptx_file: Path):
        """Test JSON data extraction with slide range."""
        data = get_ppt_data_as_json(str(temp_pptx_file), slide_range=(1, 2))
        
        assert data['success'] is True
        assert len(data['slides']) == 2

    def test_get_ppt_data_as_json_nonexistent(self):
        """Test JSON data extraction for non-existent file."""
        data = get_ppt_data_as_json("/nonexistent/file.pptx")
        
        assert data['success'] is False
        assert 'error' in data


class TestPptTextUpdate:
    """Tests for updating text in PowerPoint slides."""

    @pytest.fixture
    def editable_pptx(self, temp_dir: Path) -> Generator[Path, None, None]:
        """Fixture for creating an editable PowerPoint file."""
        file_path = temp_dir / "editable.pptx"
        
        prs = Presentation()
        
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Original Title"
        subtitle.text = "Original Subtitle"
        
        prs.save(file_path)
        
        yield file_path
        
        if file_path.exists():
            file_path.unlink()

    def test_update_text_in_slide_basic(self, editable_pptx: Path):
        """Test basic text update in a slide."""
        slide_idx = 1
        shape_idx = 0
        new_text = "Updated Title"
        
        success = update_text_in_slide(str(editable_pptx), slide_idx, shape_idx, new_text)
        
        assert success is True
        
        prs = Presentation(str(editable_pptx))
        slide = prs.slides[slide_idx - 1]
        shape = slide.shapes[shape_idx]
        
        assert shape.has_text_frame
        assert new_text in shape.text_frame.text

    def test_update_text_in_slide_invalid_slide_idx(self, editable_pptx: Path):
        """Test text update with invalid slide index."""
        success = update_text_in_slide(str(editable_pptx), 999, 0, "Test")
        
        assert success is False

    def test_update_text_in_slide_invalid_shape_idx(self, editable_pptx: Path):
        """Test text update with invalid shape index."""
        success = update_text_in_slide(str(editable_pptx), 1, 999, "Test")
        
        assert success is False

    def test_update_text_in_slide_nonexistent_file(self):
        """Test text update for non-existent file."""
        success = update_text_in_slide("/nonexistent/file.pptx", 1, 0, "Test")
        
        assert success is False

    def test_update_multiple_texts_basic(self, editable_pptx: Path):
        """Test basic multiple text updates."""
        updates = [
            {"slide_idx": 1, "shape_idx": 0, "text": "New Title"},
            {"slide_idx": 1, "shape_idx": 1, "text": "New Subtitle"},
        ]
        
        success, updated_count = update_multiple_texts(str(editable_pptx), updates)
        
        assert success is True
        assert updated_count == 2


class TestPptSlideManagement:
    """Tests for adding, deleting, and reordering slides."""

    @pytest.fixture
    def manageable_pptx(self, temp_dir: Path) -> Generator[Path, None, None]:
        """Fixture for creating a PowerPoint file with multiple slides."""
        file_path = temp_dir / "manageable.pptx"
        
        prs = Presentation()
        
        for i in range(3):
            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            title.text = f"Slide {i + 1}"
        
        prs.save(file_path)
        
        yield file_path
        
        if file_path.exists():
            file_path.unlink()

    def test_add_new_slide_basic(self, manageable_pptx: Path):
        """Test basic slide addition."""
        prs_before = Presentation(str(manageable_pptx))
        slide_count_before = len(prs_before.slides)
        del prs_before
        
        success = add_new_slide(str(manageable_pptx))
        
        assert success is True
        
        prs_after = Presentation(str(manageable_pptx))
        slide_count_after = len(prs_after.slides)
        del prs_after
        
        assert slide_count_after == slide_count_before + 1

    def test_add_new_slide_with_title(self, manageable_pptx: Path):
        """Test slide addition with title text."""
        success = add_new_slide(
            str(manageable_pptx),
            layout_idx=0,
            title_text="Test Title",
            content_text="Test Content"
        )
        
        assert success is True

    def test_add_new_slide_invalid_layout(self, manageable_pptx: Path):
        """Test slide addition with invalid layout index (should use default)."""
        success = add_new_slide(str(manageable_pptx), layout_idx=999)
        
        assert success is True

    def test_add_new_slide_nonexistent_file(self):
        """Test slide addition for non-existent file."""
        success = add_new_slide("/nonexistent/file.pptx")
        
        assert success is False

    def test_delete_slide_basic(self, manageable_pptx: Path):
        """Test basic slide deletion."""
        prs_before = Presentation(str(manageable_pptx))
        slide_count_before = len(prs_before.slides)
        del prs_before
        
        success = delete_slide(str(manageable_pptx), 2)
        
        assert success is True
        
        prs_after = Presentation(str(manageable_pptx))
        slide_count_after = len(prs_after.slides)
        del prs_after
        
        assert slide_count_after == slide_count_before - 1

    def test_delete_slide_last_slide(self, temp_dir: Path):
        """Test deleting the last slide (should fail)."""
        file_path = temp_dir / "single_slide.pptx"
        
        prs = Presentation()
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        title.text = "Single Slide"
        prs.save(file_path)
        
        success = delete_slide(str(file_path), 1)
        
        assert success is False

    def test_delete_slide_invalid_idx(self, manageable_pptx: Path):
        """Test slide deletion with invalid index."""
        success = delete_slide(str(manageable_pptx), 999)
        
        assert success is False

    def test_delete_slide_nonexistent_file(self):
        """Test slide deletion for non-existent file."""
        success = delete_slide("/nonexistent/file.pptx", 1)
        
        assert success is False

    def test_reorder_slides_basic(self, manageable_pptx: Path):
        """Test basic slide reordering."""
        prs_before = Presentation(str(manageable_pptx))
        first_slide_title_before = prs_before.slides[0].shapes.title.text
        del prs_before
        
        success = reorder_slides(str(manageable_pptx), 3, 1)
        
        assert success is True
        
        prs_after = Presentation(str(manageable_pptx))
        first_slide_title_after = prs_after.slides[0].shapes.title.text
        del prs_after
        
        assert first_slide_title_after == "Slide 3"

    def test_reorder_slides_same_position(self, manageable_pptx: Path):
        """Test reordering slide to the same position."""
        success = reorder_slides(str(manageable_pptx), 1, 1)
        
        assert success is True

    def test_reorder_slides_invalid_idx(self, manageable_pptx: Path):
        """Test reordering with invalid indices."""
        success = reorder_slides(str(manageable_pptx), 999, 1)
        
        assert success is False

    def test_reorder_slides_nonexistent_file(self):
        """Test reordering for non-existent file."""
        success = reorder_slides("/nonexistent/file.pptx", 1, 2)
        
        assert success is False
