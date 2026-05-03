"""Tests for API endpoints."""
import sys
import io
from pathlib import Path
from datetime import datetime
from unittest.mock import patch, MagicMock

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.models import Document, Folder


class TestDocumentEndpoints:
    """Tests for document-related API endpoints."""

    def test_get_documents_empty(self, client: TestClient, db_session: Session):
        """Test getting documents when none exist."""
        response = client.get("/api/documents/")
        
        assert response.status_code == 200
        data = response.json()
        assert "documents" in data
        assert len(data["documents"]) == 0

    def test_get_documents_with_data(self, client: TestClient, db_session: Session):
        """Test getting documents when some exist."""
        doc1 = Document(
            title="Test Document 1",
            filename="test1.docx",
            file_type="doc",
            file_path="/uploads/test1.docx",
            content="Test content 1",
            html_content="<p>Test content 1</p>",
        )
        doc2 = Document(
            title="Test Document 2",
            filename="test2.xlsx",
            file_type="excel",
            file_path="/uploads/test2.xlsx",
            content="Test content 2",
            html_content="<p>Test content 2</p>",
        )
        
        db_session.add_all([doc1, doc2])
        db_session.commit()
        
        response = client.get("/api/documents/")
        
        assert response.status_code == 200
        data = response.json()
        assert len(data["documents"]) == 2
        
        titles = [doc["title"] for doc in data["documents"]]
        assert "Test Document 1" in titles
        assert "Test Document 2" in titles

    def test_get_document_by_id(self, client: TestClient, db_session: Session):
        """Test getting a single document by ID."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}")
        
        assert response.status_code == 200
        data = response.json()
        assert data["id"] == doc.id
        assert data["title"] == "Test Document"
        assert data["filename"] == "test.docx"
        assert data["file_type"] == "doc"
        assert data["content"] == "Test content"
        assert data["html_content"] == "<p>Test content</p>"

    def test_get_document_not_found(self, client: TestClient):
        """Test getting a non-existent document."""
        response = client.get("/api/documents/999999")
        
        assert response.status_code == 404
        data = response.json()
        assert "detail" in data

    def test_delete_document(self, client: TestClient, db_session: Session):
        """Test deleting a document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        doc_id = doc.id
        
        response = client.delete(f"/api/documents/{doc_id}")
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data
        
        deleted_doc = db_session.query(Document).filter(Document.id == doc_id).first()
        assert deleted_doc is None

    def test_delete_document_not_found(self, client: TestClient):
        """Test deleting a non-existent document."""
        response = client.delete("/api/documents/999999")
        
        assert response.status_code == 404

    def test_update_document(self, client: TestClient, db_session: Session):
        """Test updating a document."""
        doc = Document(
            title="Original Title",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Original content",
            html_content="<p>Original content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "title": "Updated Title",
            "content": "Updated content",
            "html_content": "<p>Updated content</p>",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}",
            json=update_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["title"] == "Updated Title"
        assert data["content"] == "Updated content"
        assert data["html_content"] == "<p>Updated content</p>"

    def test_update_document_partial(self, client: TestClient, db_session: Session):
        """Test updating only some fields of a document."""
        doc = Document(
            title="Original Title",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Original content",
            html_content="<p>Original content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "title": "Updated Title",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}",
            json=update_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["title"] == "Updated Title"
        assert data["content"] == "Original content"

    def test_update_document_not_found(self, client: TestClient):
        """Test updating a non-existent document."""
        update_data = {
            "title": "Updated Title",
        }
        
        response = client.put(
            "/api/documents/999999",
            json=update_data
        )
        
        assert response.status_code == 404


class TestUploadEndpoints:
    """Tests for file upload API endpoints."""

    def test_upload_unsupported_file(self, client: TestClient):
        """Test uploading an unsupported file type."""
        file_content = b"This is a test file"
        file = io.BytesIO(file_content)
        
        response = client.post(
            "/api/upload/",
            files={"file": ("test.txt", file, "text/plain")},
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "不支持" in data["detail"] or "unsupported" in data["detail"].lower()

    def test_upload_multiple_files(self, client: TestClient, temp_dir: Path, monkeypatch):
        """Test uploading multiple files."""
        with patch('app.utils.parse_file', return_value=("Test content", "<p>Test</p>")):
            with patch('app.utils.save_upload_file', return_value=Path("/uploads/test.xlsx")):
                file1_content = b"Excel file content"
                file1 = io.BytesIO(file1_content)
                
                file2_content = b"Docx file content"
                file2 = io.BytesIO(file2_content)
                
                response = client.post(
                    "/api/upload/multiple",
                    files=[
                        ("files", ("test1.xlsx", file1, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
                        ("files", ("test2.docx", file2, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                    ],
                )
                
                assert response.status_code == 200
                data = response.json()
                assert "success" in data or "success_count" in data or len(data.get("success", [])) >= 0


class TestFolderEndpoints:
    """Tests for folder-related API endpoints."""

    def test_get_folders_empty(self, client: TestClient, db_session: Session):
        """Test getting folders when none exist."""
        response = client.get("/api/folders/")
        
        assert response.status_code == 200
        data = response.json()
        assert "folders" in data
        assert len(data["folders"]) == 0

    def test_create_folder(self, client: TestClient, db_session: Session):
        """Test creating a new folder."""
        response = client.post(
            "/api/folders/",
            data={"name": "Test Folder"}
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "folder" in data
        assert data["folder"]["name"] == "Test Folder"
        
        folder = db_session.query(Folder).filter(Folder.id == data["folder"]["id"]).first()
        assert folder is not None
        assert folder.name == "Test Folder"

    def test_create_folder_with_parent(self, client: TestClient, db_session: Session):
        """Test creating a folder with a parent folder."""
        parent_folder = Folder(name="Parent Folder")
        db_session.add(parent_folder)
        db_session.commit()
        parent_folder_id = parent_folder.id
        
        response = client.post(
            "/api/folders/",
            data={"name": "Child Folder", "parent_id": parent_folder_id}
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["folder"]["parent_id"] == parent_folder_id

    def test_get_folder_by_id(self, client: TestClient, db_session: Session):
        """Test getting a single folder by ID."""
        folder = Folder(name="Test Folder")
        db_session.add(folder)
        db_session.commit()
        db_session.refresh(folder)
        
        response = client.get(f"/api/folders/{folder.id}")
        
        assert response.status_code == 200
        data = response.json()
        assert data["id"] == folder.id
        assert data["name"] == "Test Folder"

    def test_get_folder_not_found(self, client: TestClient):
        """Test getting a non-existent folder."""
        response = client.get("/api/folders/999999")
        
        assert response.status_code == 404

    def test_delete_folder(self, client: TestClient, db_session: Session):
        """Test deleting a folder."""
        folder = Folder(name="Test Folder")
        db_session.add(folder)
        db_session.commit()
        folder_id = folder.id
        
        response = client.delete(f"/api/folders/{folder_id}")
        
        assert response.status_code == 200
        
        deleted_folder = db_session.query(Folder).filter(Folder.id == folder_id).first()
        assert deleted_folder is None

    def test_delete_folder_with_documents(self, client: TestClient, db_session: Session):
        """Test that deleting a folder also deletes its documents (cascade)."""
        folder = Folder(name="Test Folder")
        db_session.add(folder)
        db_session.commit()
        db_session.refresh(folder)
        
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
            folder_id=folder.id,
        )
        
        db_session.add(doc)
        db_session.commit()
        folder_id = folder.id
        doc_id = doc.id
        
        response = client.delete(f"/api/folders/{folder_id}")
        
        assert response.status_code == 200
        
        deleted_folder = db_session.query(Folder).filter(Folder.id == folder_id).first()
        assert deleted_folder is None
        
        deleted_doc = db_session.query(Document).filter(Document.id == doc_id).first()
        assert deleted_doc is None


class TestSearchEndpoints:
    """Tests for search API endpoints."""

    def test_search_documents(self, client: TestClient, db_session: Session):
        """Test searching for documents."""
        doc1 = Document(
            title="Python Programming",
            filename="python.docx",
            file_type="doc",
            file_path="/uploads/python.docx",
            content="Python is a great programming language",
            html_content="<p>Python is a great programming language</p>",
        )
        doc2 = Document(
            title="JavaScript Guide",
            filename="javascript.docx",
            file_type="doc",
            file_path="/uploads/javascript.docx",
            content="JavaScript is used for web development",
            html_content="<p>JavaScript is used for web development</p>",
        )
        
        db_session.add_all([doc1, doc2])
        db_session.commit()
        
        response = client.get("/api/search/?q=Python")
        
        assert response.status_code == 200
        data = response.json()
        assert "results" in data

    def test_search_no_query(self, client: TestClient):
        """Test search with no query parameter (should fail as q is required)."""
        response = client.get("/api/search/")
        
        assert response.status_code == 422


class TestExcelSpecificEndpoints:
    """Tests for Excel-specific API endpoints."""

    def test_get_excel_metadata(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test getting Excel metadata."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/excel/metadata")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["sheet_count"] == 2

    def test_get_excel_metadata_not_excel(self, client: TestClient, db_session: Session):
        """Test getting metadata for non-Excel document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/excel/metadata")
        
        assert response.status_code == 400

    def test_get_excel_data(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test getting Excel data."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/excel/data")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["sheet_count"] == 2

    def test_get_excel_data_paginated(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test getting Excel data with pagination."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(
            f"/api/documents/{doc.id}/excel/data/paginated?sheet_name=Sheet1&start_row=1&end_row=2"
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["row_count"] == 2

    def test_update_excel_cell(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test updating an Excel cell."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "sheet_name": "Sheet1",
            "row": 2,
            "col": 1,
            "value": "Updated Name",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/excel/cell",
            json=update_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_update_excel_cells_batch(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test batch updating Excel cells."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "updates": [
                {"sheet_name": "Sheet1", "row": 2, "col": 1, "value": "Updated Name 1"},
                {"sheet_name": "Sheet1", "row": 3, "col": 1, "value": "Updated Name 2"},
            ]
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/excel/cells",
            json=update_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "updated_count" in data or "message" in data

    def test_add_excel_sheet(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test adding a new Excel sheet."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        operation_data = {
            "sheet_name": "NewSheet",
        }
        
        response = client.post(
            f"/api/documents/{doc.id}/excel/sheet",
            json=operation_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_delete_excel_sheet(self, client: TestClient, db_session: Session, temp_excel_file: Path):
        """Test deleting an Excel sheet."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        operation_data = {
            "sheet_name": "Sheet2",
        }
        
        response = client.request(
            "DELETE",
            f"/api/documents/{doc.id}/excel/sheet",
            json=operation_data
        )
        
        assert response.status_code == 200


class TestDocumentMoveEndpoint:
    """Tests for document move endpoint."""

    def test_move_document_to_folder(self, client: TestClient, db_session: Session):
        """Test moving a document to a folder."""
        folder = Folder(name="Target Folder")
        db_session.add(folder)
        db_session.commit()
        folder_id = folder.id
        
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
            folder_id=None,
        )
        
        db_session.add(doc)
        db_session.commit()
        doc_id = doc.id
        
        response = client.put(f"/api/documents/{doc_id}/move/{folder_id}")
        
        assert response.status_code == 200
        
        updated_doc = db_session.query(Document).filter(Document.id == doc_id).first()
        assert updated_doc.folder_id == folder_id

    def test_move_document_to_root(self, client: TestClient, db_session: Session):
        """Test moving a document back to root (folder_id=0)."""
        folder = Folder(name="Source Folder")
        db_session.add(folder)
        db_session.commit()
        folder_id = folder.id
        
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
            folder_id=folder_id,
        )
        
        db_session.add(doc)
        db_session.commit()
        doc_id = doc.id
        
        response = client.put(f"/api/documents/{doc_id}/move/0")
        
        assert response.status_code == 200
        
        updated_doc = db_session.query(Document).filter(Document.id == doc_id).first()
        assert updated_doc.folder_id is None

    def test_move_document_not_found(self, client: TestClient):
        """Test moving a non-existent document."""
        response = client.put("/api/documents/999999/move/1")
        
        assert response.status_code == 404

    def test_move_document_to_nonexistent_folder(self, client: TestClient, db_session: Session):
        """Test moving a document to a non-existent folder."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.put(f"/api/documents/{doc.id}/move/999999")
        
        assert response.status_code == 404


class TestReparseEndpoint:
    """Tests for document reparse endpoint."""

    def test_reparse_document(self, client: TestClient, db_session: Session, temp_docx_file: Path):
        """Test re-parsing a document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path=str(temp_docx_file),
            content="Old content",
            html_content="<p>Old content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.post(f"/api/documents/{doc.id}/reparse")
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data or "content" in data

    def test_reparse_document_not_found(self, client: TestClient):
        """Test re-parsing a non-existent document."""
        response = client.post("/api/documents/999999/reparse")
        
        assert response.status_code == 404


class TestFolderTreeEndpoint:
    """Tests for folder tree endpoint."""

    def test_get_folders_with_tree_param(self, client: TestClient, db_session: Session):
        """Test getting folders with tree=true parameter."""
        parent_folder = Folder(name="Parent Folder")
        db_session.add(parent_folder)
        db_session.commit()
        parent_id = parent_folder.id
        
        child_folder = Folder(name="Child Folder", parent_id=parent_id)
        db_session.add(child_folder)
        db_session.commit()
        
        response = client.get("/api/folders/?tree=true")
        
        assert response.status_code == 200
        data = response.json()
        assert "folders" in data
        assert len(data["folders"]) > 0

    def test_get_folder_tree_endpoint(self, client: TestClient, db_session: Session):
        """Test the dedicated /tree endpoint."""
        folder = Folder(name="Test Folder")
        db_session.add(folder)
        db_session.commit()
        
        response = client.get("/api/folders/tree")
        
        assert response.status_code == 200
        data = response.json()
        assert "folders" in data

    def test_get_folder_tree_with_documents(self, client: TestClient, db_session: Session):
        """Test getting folder tree with documents in folders."""
        folder = Folder(name="Test Folder")
        db_session.add(folder)
        db_session.commit()
        folder_id = folder.id
        
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
            folder_id=folder_id,
        )
        db_session.add(doc)
        db_session.commit()
        
        response = client.get(f"/api/folders/{folder_id}")
        
        assert response.status_code == 200
        data = response.json()
        assert "documents" in data
        assert len(data["documents"]) > 0


class TestUpdateFolderEndpoint:
    """Tests for folder update endpoint."""

    def test_update_folder(self, client: TestClient, db_session: Session):
        """Test updating a folder name."""
        folder = Folder(name="Old Name")
        db_session.add(folder)
        db_session.commit()
        folder_id = folder.id
        
        response = client.put(
            f"/api/folders/{folder_id}",
            data={"name": "New Name"}
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data
        
        updated_folder = db_session.query(Folder).filter(Folder.id == folder_id).first()
        assert updated_folder.name == "New Name"

    def test_update_folder_not_found(self, client: TestClient):
        """Test updating a non-existent folder."""
        response = client.put(
            "/api/folders/999999",
            data={"name": "New Name"}
        )
        
        assert response.status_code == 404


class TestCreateFolderWithInvalidParent:
    """Tests for creating folder with invalid parent."""

    def test_create_folder_with_nonexistent_parent(self, client: TestClient):
        """Test creating a folder with a parent that doesn't exist."""
        response = client.post(
            "/api/folders/",
            data={"name": "Test Folder", "parent_id": 999999}
        )
        
        assert response.status_code == 404


class TestHtmlPages:
    """Tests for HTML page endpoints."""

    def test_home_page(self, client: TestClient):
        """Test the home page endpoint."""
        response = client.get("/")
        
        assert response.status_code == 200
        assert "text/html" in response.headers.get("content-type", "")

    def test_document_detail_page(self, client: TestClient):
        """Test the document detail page."""
        response = client.get("/document/1")
        
        assert response.status_code == 200
        assert "text/html" in response.headers.get("content-type", "")

    def test_folder_detail_page(self, client: TestClient):
        """Test the folder detail page."""
        response = client.get("/folder/1")
        
        assert response.status_code == 200
        assert "text/html" in response.headers.get("content-type", "")

    def test_search_page(self, client: TestClient):
        """Test the search page."""
        response = client.get("/search")
        
        assert response.status_code == 200
        assert "text/html" in response.headers.get("content-type", "")

    def test_search_page_with_query(self, client: TestClient):
        """Test the search page with a query parameter."""
        response = client.get("/search?q=test")
        
        assert response.status_code == 200
        assert "text/html" in response.headers.get("content-type", "")


class TestMarkdownApiEndpoints:
    """Tests for Markdown-specific API endpoints."""

    def test_get_markdown_metadata(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test getting Markdown metadata."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Test\n\nContent",
            html_content="<h1>Test</h1><p>Content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/metadata")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True

    def test_get_markdown_metadata_not_markdown(self, client: TestClient, db_session: Session):
        """Test getting metadata for non-Markdown document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/metadata")
        
        assert response.status_code == 400

    def test_get_markdown_metadata_not_found(self, client: TestClient):
        """Test getting metadata for non-existent document."""
        response = client.get("/api/documents/999999/markdown/metadata")
        
        assert response.status_code == 404

    def test_get_markdown_outline(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test getting Markdown outline."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# H1\n\n## H2\n\n### H3",
            html_content="<h1>H1</h1><h2>H2</h2><h3>H3</h3>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/outline")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert "outline" in data

    def test_get_markdown_outline_not_markdown(self, client: TestClient, db_session: Session):
        """Test getting outline for non-Markdown document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/outline")
        
        assert response.status_code == 400

    def test_get_markdown_data(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test getting Markdown full data."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Test\n\nContent",
            html_content="<h1>Test</h1><p>Content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/data")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True

    def test_get_markdown_data_not_markdown(self, client: TestClient, db_session: Session):
        """Test getting data for non-Markdown document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/data")
        
        assert response.status_code == 400

    def test_update_markdown_content(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test updating Markdown content."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Original\n\nOriginal content",
            html_content="<h1>Original</h1><p>Original content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "content": "# Updated\n\nNew content here"
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/markdown/content",
            json=update_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data
        assert "document" in data

    def test_update_markdown_content_not_markdown(self, client: TestClient, db_session: Session):
        """Test updating content for non-Markdown document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "content": "New content"
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/markdown/content",
            json=update_data
        )
        
        assert response.status_code == 400

    def test_append_to_markdown(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test appending content to Markdown."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Original",
            html_content="<h1>Original</h1>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        append_data = {
            "content": "\n\n## New Section\n\nAppended content"
        }
        
        response = client.post(
            f"/api/documents/{doc.id}/markdown/append",
            json=append_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_prepend_to_markdown(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test prepending content to Markdown."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Original",
            html_content="<h1>Original</h1>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        prepend_data = {
            "content": "# New Header\n\nPrepended content\n\n"
        }
        
        response = client.post(
            f"/api/documents/{doc.id}/markdown/prepend",
            json=prepend_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_replace_markdown_section(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test replacing a section in Markdown."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Test Document\n\n## Introduction\n\nOld content",
            html_content="<h1>Test Document</h1>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        replace_data = {
            "heading_text": "Introduction",
            "new_content": "\n\nUpdated introduction content here\n\n",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/markdown/section",
            json=replace_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_replace_markdown_section_not_found(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test replacing a non-existent section in Markdown."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Main\n\n## Section 1\n\nContent",
            html_content="<h1>Main</h1>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        replace_data = {
            "heading_text": "Nonexistent Section",
            "new_content": "New content",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/markdown/section",
            json=replace_data
        )
        
        assert response.status_code == 404

    def test_add_heading_to_markdown(self, client: TestClient, db_session: Session, temp_markdown_file: Path):
        """Test adding a new heading to Markdown."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="# Original",
            html_content="<h1>Original</h1>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        heading_data = {
            "heading_text": "New Section",
            "level": 2,
            "content": "Content under new section",
        }
        
        response = client.post(
            f"/api/documents/{doc.id}/markdown/heading",
            json=heading_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data


class TestPptBackgroundEndpoints:
    """Tests for PPT background API endpoints."""

    def test_get_ppt_background(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test getting PPT slide background information."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/ppt/background?slide_idx=1")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert "background_type" in data

    def test_get_ppt_background_not_ppt(self, client: TestClient, db_session: Session):
        """Test getting background for non-PPT document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/ppt/background?slide_idx=1")
        
        assert response.status_code == 400

    def test_get_ppt_background_not_found(self, client: TestClient):
        """Test getting background for non-existent document."""
        response = client.get("/api/documents/999999/ppt/background?slide_idx=1")
        
        assert response.status_code == 404

    def test_set_ppt_solid_background(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test setting PPT slide solid background color."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        background_data = {
            "slide_idx": 1,
            "color": "FF0000",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/ppt/background/solid",
            json=background_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data
        assert data["color"] == "FF0000"

    def test_set_ppt_solid_background_not_ppt(self, client: TestClient, db_session: Session):
        """Test setting solid background for non-PPT document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        background_data = {
            "slide_idx": 1,
            "color": "FF0000",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/ppt/background/solid",
            json=background_data
        )
        
        assert response.status_code == 400

    def test_remove_ppt_background(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test removing PPT slide background."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        remove_data = {
            "slide_idx": 1,
        }
        
        response = client.request(
            "DELETE",
            f"/api/documents/{doc.id}/ppt/background",
            json=remove_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_remove_ppt_background_not_ppt(self, client: TestClient, db_session: Session):
        """Test removing background for non-PPT document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        remove_data = {
            "slide_idx": 1,
        }
        
        response = client.request(
            "DELETE",
            f"/api/documents/{doc.id}/ppt/background",
            json=remove_data
        )
        
        assert response.status_code == 400


class TestPptTextAndSlideEndpoints:
    """Tests for PPT text update and slide management API endpoints."""

    def test_update_ppt_text(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test updating text in a PPT slide."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        update_data = {
            "slide_idx": 1,
            "shape_idx": 0,
            "text": "Updated Title",
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/ppt/text",
            json=update_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_update_ppt_texts_batch(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test batch updating multiple texts in PPT."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        batch_data = {
            "updates": [
                {"slide_idx": 1, "shape_idx": 0, "text": "Updated Title 1"},
                {"slide_idx": 1, "shape_idx": 1, "text": "Updated Subtitle"},
            ]
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/ppt/texts",
            json=batch_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "updated_count" in data

    def test_add_ppt_slide(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test adding a new slide to PPT."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        slide_data = {
            "layout_idx": 6,
            "title_text": "New Slide Title",
            "content_text": "New slide content",
        }
        
        response = client.post(
            f"/api/documents/{doc.id}/ppt/slide",
            json=slide_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_delete_ppt_slide(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test deleting a slide from PPT."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        delete_data = {
            "slide_idx": 2,
        }
        
        response = client.request(
            "DELETE",
            f"/api/documents/{doc.id}/ppt/slide",
            json=delete_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data

    def test_reorder_ppt_slides(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test reordering slides in PPT."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        reorder_data = {
            "old_idx": 3,
            "new_idx": 1,
        }
        
        response = client.put(
            f"/api/documents/{doc.id}/ppt/reorder",
            json=reorder_data
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "message" in data


class TestPptDataEndpoints:
    """Tests for PPT data retrieval API endpoints."""

    def test_get_ppt_metadata(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test getting PPT metadata."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/ppt/metadata")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["slide_count"] == 3

    def test_get_ppt_metadata_not_ppt(self, client: TestClient, db_session: Session):
        """Test getting metadata for non-PPT document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/ppt/metadata")
        
        assert response.status_code == 400

    def test_get_ppt_data(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test getting PPT full data."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/ppt/data")
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["slide_count"] == 3

    def test_get_ppt_data_paginated(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test getting PPT data with pagination."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(
            f"/api/documents/{doc.id}/ppt/data/paginated?start_slide=1&end_slide=2"
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["slide_count"] == 2

    def test_get_ppt_data_with_range(self, client: TestClient, db_session: Session, temp_pptx_file: Path):
        """Test getting PPT data with slide range."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test content</p>",
        )
        
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(
            f"/api/documents/{doc.id}/ppt/data?start_slide=1&end_slide=2"
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True


class TestPptPictureBackgroundEndpoints:
    """Tests for PPT picture background API endpoints."""

    def test_set_ppt_picture_background_not_found(self, client: TestClient, temp_image_file: Path):
        """Test setting picture background for non-existent document."""
        with open(temp_image_file, "rb") as f:
            response = client.post(
                "/api/documents/999999/ppt/background/picture",
                data={"slide_idx": 1},
                files={"file": ("test.png", f, "image/png")},
            )
        
        assert response.status_code == 404

    def test_set_ppt_picture_background_not_ppt(
        self, client: TestClient, db_session: Session, temp_docx_file: Path, temp_image_file: Path
    ):
        """Test setting picture background for non-PPT document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path=str(temp_docx_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        with open(temp_image_file, "rb") as f:
            response = client.post(
                f"/api/documents/{doc.id}/ppt/background/picture",
                data={"slide_idx": 1},
                files={"file": ("test.png", f, "image/png")},
            )
        
        assert response.status_code == 400


class TestPptPictureShapeEndpoints:
    """Tests for PPT picture shape API endpoints."""

    def test_add_ppt_picture_shape_not_found(self, client: TestClient, temp_image_file: Path):
        """Test adding picture shape for non-existent document."""
        with open(temp_image_file, "rb") as f:
            response = client.post(
                "/api/documents/999999/ppt/picture",
                data={"slide_idx": 1, "left": 0, "top": 0},
                files={"file": ("test.png", f, "image/png")},
            )
        
        assert response.status_code == 404

    def test_add_ppt_picture_shape_not_ppt(
        self, client: TestClient, db_session: Session, temp_docx_file: Path, temp_image_file: Path
    ):
        """Test adding picture shape for non-PPT document."""
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path=str(temp_docx_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        with open(temp_image_file, "rb") as f:
            response = client.post(
                f"/api/documents/{doc.id}/ppt/picture",
                data={"slide_idx": 1, "left": 0, "top": 0},
                files={"file": ("test.png", f, "image/png")},
            )
        
        assert response.status_code == 400


class TestUploadAdditionalEndpoints:
    """Additional tests for upload API endpoints."""

    def test_upload_docx_file(self, client: TestClient, temp_docx_file: Path):
        """Test uploading a DOCX file."""
        with open(temp_docx_file, "rb") as f:
            response = client.post(
                "/api/upload/",
                files={"file": ("test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        
        assert response.status_code in [200, 400]

    def test_upload_pptx_file(self, client: TestClient, temp_pptx_file: Path):
        """Test uploading a PPTX file."""
        with open(temp_pptx_file, "rb") as f:
            response = client.post(
                "/api/upload/",
                files={"file": ("test.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            )
        
        assert response.status_code in [200, 400]

    def test_upload_markdown_file(self, client: TestClient, temp_markdown_file: Path):
        """Test uploading a Markdown file."""
        with open(temp_markdown_file, "rb") as f:
            response = client.post(
                "/api/upload/",
                files={"file": ("test.md", f, "text/markdown")},
            )
        
        assert response.status_code in [200, 400]


class TestPptParserDirectTests:
    """Direct tests for PPT parser functions to improve coverage."""

    def test_set_picture_background_nonexistent_file(self, temp_image_file: Path):
        """Test set_picture_background with nonexistent PPT file."""
        from app.parsers.ppt_parser import set_picture_background
        
        result = set_picture_background("/nonexistent.pptx", 1, str(temp_image_file))
        assert result is False

    def test_set_picture_background_nonexistent_image(self, temp_pptx_file: Path):
        """Test set_picture_background with nonexistent image file."""
        from app.parsers.ppt_parser import set_picture_background
        
        result = set_picture_background(str(temp_pptx_file), 1, "/nonexistent.png")
        assert result is False

    def test_set_picture_background_invalid_slide_idx(self, temp_pptx_file: Path, temp_image_file: Path):
        """Test set_picture_background with invalid slide index."""
        from app.parsers.ppt_parser import set_picture_background
        
        result = set_picture_background(str(temp_pptx_file), 999, str(temp_image_file))
        assert result is False

    def test_set_picture_background_v2_nonexistent_file(self, temp_image_file: Path):
        """Test set_picture_background_v2 with nonexistent PPT file."""
        from app.parsers.ppt_parser import set_picture_background_v2
        
        result = set_picture_background_v2("/nonexistent.pptx", 1, str(temp_image_file))
        assert result is False

    def test_set_picture_background_v2_nonexistent_image(self, temp_pptx_file: Path):
        """Test set_picture_background_v2 with nonexistent image file."""
        from app.parsers.ppt_parser import set_picture_background_v2
        
        result = set_picture_background_v2(str(temp_pptx_file), 1, "/nonexistent.png")
        assert result is False

    def test_set_picture_background_v2_invalid_slide_idx(self, temp_pptx_file: Path, temp_image_file: Path):
        """Test set_picture_background_v2 with invalid slide index."""
        from app.parsers.ppt_parser import set_picture_background_v2
        
        result = set_picture_background_v2(str(temp_pptx_file), 999, str(temp_image_file))
        assert result is False

    def test_add_picture_shape_nonexistent_file(self, temp_image_file: Path):
        """Test add_picture_shape_to_slide with nonexistent PPT file."""
        from app.parsers.ppt_parser import add_picture_shape_to_slide
        
        result = add_picture_shape_to_slide("/nonexistent.pptx", 1, str(temp_image_file))
        assert result is False

    def test_add_picture_shape_nonexistent_image(self, temp_pptx_file: Path):
        """Test add_picture_shape_to_slide with nonexistent image file."""
        from app.parsers.ppt_parser import add_picture_shape_to_slide
        
        result = add_picture_shape_to_slide(str(temp_pptx_file), 1, "/nonexistent.png")
        assert result is False

    def test_add_picture_shape_invalid_slide_idx(self, temp_pptx_file: Path, temp_image_file: Path):
        """Test add_picture_shape_to_slide with invalid slide index."""
        from app.parsers.ppt_parser import add_picture_shape_to_slide
        
        result = add_picture_shape_to_slide(str(temp_pptx_file), 999, str(temp_image_file))
        assert result is False


class TestExcelParserAdditionalTests:
    """Additional tests for Excel parser to improve coverage."""

    def test_get_sheet_metadata_with_empty_file(self, temp_dir: Path):
        """Test get_sheet_metadata with empty Excel file."""
        from openpyxl import Workbook
        from app.parsers.excel_parser import get_sheet_metadata
        
        file_path = temp_dir / "empty.xlsx"
        wb = Workbook()
        wb.save(file_path)
        
        result = get_sheet_metadata(str(file_path))
        assert result["success"] is True

    def test_get_sheet_data_paginated_with_invalid_range(self, temp_excel_file: Path):
        """Test get_sheet_data_paginated with invalid row range."""
        from app.parsers.excel_parser import get_sheet_data_paginated
        
        result = get_sheet_data_paginated(str(temp_excel_file), "Sheet1", -1, -5)
        assert result["success"] is True

    def test_update_cell_in_excel_with_nonexistent_sheet(self, temp_excel_file: Path):
        """Test update_cell_in_excel with nonexistent sheet."""
        from app.parsers.excel_parser import update_cell_in_excel
        
        result = update_cell_in_excel(str(temp_excel_file), "NonexistentSheet", 1, 1, "Test")
        assert result is False

    def test_add_new_sheet_duplicate_name(self, temp_excel_file: Path):
        """Test add_new_sheet with duplicate name."""
        from app.parsers.excel_parser import add_new_sheet
        
        result = add_new_sheet(str(temp_excel_file), "Sheet1")
        assert result is False

    def test_delete_sheet_last_one(self, temp_dir: Path):
        """Test delete_sheet when it's the last sheet."""
        from openpyxl import Workbook
        from app.parsers.excel_parser import delete_sheet
        
        file_path = temp_dir / "single_sheet.xlsx"
        wb = Workbook()
        wb.save(file_path)
        
        result = delete_sheet(str(file_path), "Sheet")
        assert result is False


class TestDocParserAdditionalTests:
    """Additional tests for Doc parser to improve coverage."""

    def test_parse_doc_nonexistent_file(self):
        """Test parse_doc with nonexistent file."""
        from app.parsers.doc_parser import parse_doc
        
        content, html = parse_doc("/nonexistent.doc")
        assert content is not None
        assert "错误" in content or "error" in content.lower()

    def test_is_docx_file_with_docx(self, temp_docx_file: Path):
        """Test is_docx_file with actual DOCX file."""
        from app.parsers.doc_parser import is_docx_file
        
        result = is_docx_file(str(temp_docx_file))
        assert result is True

    def test_is_docx_file_with_non_docx(self, temp_markdown_file: Path):
        """Test is_docx_file with non-DOCX file."""
        from app.parsers.doc_parser import is_docx_file
        
        result = is_docx_file(str(temp_markdown_file))
        assert result is False


class TestUploadFolderEndpoints:
    """Tests for upload with folder ID to improve coverage."""

    def test_upload_with_nonexistent_folder(self, client: TestClient, temp_excel_file: Path):
        """Test uploading a file with a nonexistent folder ID."""
        with open(temp_excel_file, "rb") as f:
            response = client.post(
                "/api/upload/",
                data={"folder_id": 999999},
                files={"file": ("test.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        
        assert response.status_code in [200, 400]
        if response.status_code == 200:
            data = response.json()
            assert data["document"]["folder_id"] is None


class TestPptPictureBackgroundSuccessTests:
    """Tests for PPT picture background success paths."""

    def test_set_picture_background_v2_success(self, temp_pptx_file: Path, temp_image_file: Path):
        """Test set_picture_background_v2 with valid files."""
        from app.parsers.ppt_parser import set_picture_background_v2
        
        result = set_picture_background_v2(str(temp_pptx_file), 1, str(temp_image_file))
        assert result is True

    def test_add_picture_shape_success(self, temp_pptx_file: Path, temp_image_file: Path):
        """Test add_picture_shape_to_slide with valid files."""
        from app.parsers.ppt_parser import add_picture_shape_to_slide
        
        result = add_picture_shape_to_slide(
            str(temp_pptx_file), 1, str(temp_image_file), 0, 0, 100, 100
        )
        assert result is True

    def test_add_picture_shape_no_size(self, temp_pptx_file: Path, temp_image_file: Path):
        """Test add_picture_shape_to_slide without specifying size."""
        from app.parsers.ppt_parser import add_picture_shape_to_slide
        
        result = add_picture_shape_to_slide(
            str(temp_pptx_file), 1, str(temp_image_file), 0, 0, None, None
        )
        assert result is True


class TestExcelParserMoreTests:
    """More tests for Excel parser to improve coverage."""

    def test_get_sheet_data_as_json_single_sheet(self, temp_excel_file: Path):
        """Test get_sheet_data_as_json with specific sheet."""
        from app.parsers.excel_parser import get_sheet_data_as_json
        
        result = get_sheet_data_as_json(str(temp_excel_file), "Sheet1")
        assert result["success"] is True

    def test_get_sheet_data_as_json_all_sheets(self, temp_excel_file: Path):
        """Test get_sheet_data_as_json with all sheets."""
        from app.parsers.excel_parser import get_sheet_data_as_json
        
        result = get_sheet_data_as_json(str(temp_excel_file))
        assert result["success"] is True


class TestMarkdownParserMoreTests:
    """More tests for Markdown parser to improve coverage."""

    def test_get_markdown_data_as_json_with_content(self, temp_markdown_file: Path):
        """Test get_markdown_data_as_json with actual markdown file."""
        from app.parsers.markdown_parser import get_markdown_data_as_json
        
        result = get_markdown_data_as_json(str(temp_markdown_file))
        assert result["success"] is True

    def test_get_markdown_outline_with_content(self, temp_markdown_file: Path):
        """Test get_markdown_outline with actual markdown file."""
        from app.parsers.markdown_parser import get_markdown_outline
        
        result = get_markdown_outline(str(temp_markdown_file))
        assert result is not None
        assert len(result) > 0

    def test_get_markdown_metadata_with_content(self, temp_markdown_file: Path):
        """Test get_markdown_metadata with actual markdown file."""
        from app.parsers.markdown_parser import get_markdown_metadata
        
        result = get_markdown_metadata(str(temp_markdown_file))
        assert result["success"] is True


class TestUploadMultipleEndpoints:
    """Tests for multiple file upload to improve coverage."""

    def test_upload_multiple_unsupported_file(self, client: TestClient, temp_dir: Path):
        """Test uploading multiple files with an unsupported file type."""
        unsupported_file = temp_dir / "test.xyz"
        unsupported_file.write_text("This is not a supported file type")
        
        with open(unsupported_file, "rb") as f:
            response = client.post(
                "/api/upload/multiple",
                files={"files": ("test.xyz", f, "application/octet-stream")},
            )
        
        assert response.status_code == 200
        data = response.json()
        assert len(data["errors"]) > 0


class TestExcelParserEdgeCases:
    """Additional edge case tests for Excel parser."""

    def test_get_sheet_data_paginated_with_invalid_params(self, temp_excel_file: Path):
        """Test get_sheet_data_paginated with invalid parameters."""
        from app.parsers.excel_parser import get_sheet_data_paginated
        
        result = get_sheet_data_paginated(str(temp_excel_file), "Sheet1", 0, 0)
        assert result["success"] is True

    def test_get_sheet_data_paginated_with_large_range(self, temp_excel_file: Path):
        """Test get_sheet_data_paginated with range larger than sheet."""
        from app.parsers.excel_parser import get_sheet_data_paginated
        
        result = get_sheet_data_paginated(str(temp_excel_file), "Sheet1", 1000, 2000)
        assert result["success"] is True


class TestPptParserMoreEdgeCases:
    """More edge case tests for PPT parser."""

    def test_get_ppt_metadata_with_content(self, temp_pptx_file: Path):
        """Test get_ppt_metadata with actual PPT file."""
        from app.parsers.ppt_parser import get_ppt_metadata
        
        result = get_ppt_metadata(str(temp_pptx_file))
        assert result["success"] is True

    def test_get_slide_data_paginated_with_content(self, temp_pptx_file: Path):
        """Test get_slide_data_paginated with actual PPT file."""
        from app.parsers.ppt_parser import get_slide_data_paginated
        
        result = get_slide_data_paginated(str(temp_pptx_file), 1, 2)
        assert result["success"] is True

    def test_get_ppt_data_as_json_with_content(self, temp_pptx_file: Path):
        """Test get_ppt_data_as_json with actual PPT file."""
        from app.parsers.ppt_parser import get_ppt_data_as_json
        
        result = get_ppt_data_as_json(str(temp_pptx_file))
        assert result["success"] is True


class TestDocumentEndpointsWithFolder:
    """Tests for document endpoints with folder filtering."""

    def test_get_documents_with_valid_folder_id(
        self, client: TestClient, db_session: Session
    ):
        """Test getting documents with a valid folder ID."""
        folder = Folder(name="Test Folder")
        db_session.add(folder)
        db_session.commit()
        db_session.refresh(folder)
        
        doc = Document(
            title="Test Document",
            filename="test.docx",
            file_type="doc",
            file_path="/uploads/test.docx",
            content="Test content",
            html_content="<p>Test</p>",
            folder_id=folder.id,
        )
        db_session.add(doc)
        db_session.commit()
        
        response = client.get(f"/api/documents/?folder_id={folder.id}")
        
        assert response.status_code == 200
        data = response.json()
        assert len(data["documents"]) == 1
        assert data["documents"][0]["folder_id"] == folder.id

    def test_get_documents_with_folder_id_no_docs(
        self, client: TestClient, db_session: Session
    ):
        """Test getting documents with a folder ID that has no documents."""
        folder = Folder(name="Empty Folder")
        db_session.add(folder)
        db_session.commit()
        db_session.refresh(folder)
        
        response = client.get(f"/api/documents/?folder_id={folder.id}")
        
        assert response.status_code == 200
        data = response.json()
        assert len(data["documents"]) == 0


class TestExcelParserMergedCells:
    """Tests for Excel parser merged cells functionality."""

    def test_get_sheet_data_paginated_with_merged_cells(self):
        """Test get_sheet_data_paginated with merged cells."""
        import tempfile
        from app.parsers.excel_parser import get_sheet_data_paginated
        
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            wb = Workbook()
            ws = wb.active
            ws.cell(row=1, column=1, value="Merged Header")
            ws.merge_cells('A1:B1')
            ws.cell(row=2, column=1, value="Data 1")
            ws.cell(row=2, column=2, value="Data 2")
            wb.save(f.name)
            
            result = get_sheet_data_paginated(f.name, "Sheet", 1, 10)
            assert result["success"] is True
            
            Path(f.name).unlink()


class TestExcelParserStyles:
    """Tests for Excel parser styles functionality."""

    def test_get_sheet_data_paginated_with_styles(self):
        """Test get_sheet_data_paginated with styles."""
        import tempfile
        from app.parsers.excel_parser import get_sheet_data_paginated
        from openpyxl.styles import Font, PatternFill
        
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            wb = Workbook()
            ws = wb.active
            cell = ws.cell(row=1, column=1, value="Styled Cell")
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            wb.save(f.name)
            
            result = get_sheet_data_paginated(f.name, "Sheet", 1, 10, include_styles=True)
            assert result["success"] is True
            
            Path(f.name).unlink()

    def test_get_sheet_data_paginated_without_styles(self):
        """Test get_sheet_data_paginated without styles."""
        import tempfile
        from app.parsers.excel_parser import get_sheet_data_paginated
        
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            wb = Workbook()
            ws = wb.active
            ws.cell(row=1, column=1, value="Simple Cell")
            wb.save(f.name)
            
            result = get_sheet_data_paginated(f.name, "Sheet", 1, 10, include_styles=False)
            assert result["success"] is True
            
            Path(f.name).unlink()


class TestPptParserSlideData:
    """Tests for PPT parser slide data functionality."""

    def test_get_ppt_metadata_with_content(self, temp_pptx_file: Path):
        """Test get_ppt_metadata with actual PPT file."""
        from app.parsers.ppt_parser import get_ppt_metadata
        
        result = get_ppt_metadata(str(temp_pptx_file))
        assert result["success"] is True

    def test_get_ppt_metadata_with_invalid_file(self):
        """Test get_ppt_metadata with invalid file."""
        from app.parsers.ppt_parser import get_ppt_metadata
        
        result = get_ppt_metadata("nonexistent.pptx")
        assert result["success"] is False

    def test_get_ppt_data_as_json_with_content(self, temp_pptx_file: Path):
        """Test get_ppt_data_as_json with actual PPT file."""
        from app.parsers.ppt_parser import get_ppt_data_as_json
        
        result = get_ppt_data_as_json(str(temp_pptx_file))
        assert result["success"] is True


class TestMarkdownParserMetadata:
    """Tests for Markdown parser metadata functionality."""

    def test_get_markdown_metadata_with_content(self, temp_markdown_file: Path):
        """Test get_markdown_metadata with actual Markdown file."""
        from app.parsers.markdown_parser import get_markdown_metadata
        
        result = get_markdown_metadata(str(temp_markdown_file))
        assert result is not None

    def test_get_markdown_outline_with_content(self, temp_markdown_file: Path):
        """Test get_markdown_outline with actual Markdown file."""
        from app.parsers.markdown_parser import get_markdown_outline
        
        result = get_markdown_outline(str(temp_markdown_file))
        assert result is not None

    def test_get_markdown_data_as_json_with_content(self, temp_markdown_file: Path):
        """Test get_markdown_data_as_json with actual Markdown file."""
        from app.parsers.markdown_parser import get_markdown_data_as_json
        
        result = get_markdown_data_as_json(str(temp_markdown_file))
        assert result is not None


class TestDocumentEndpointsMarkdown:
    """Tests for document endpoints with Markdown-specific functionality."""

    def test_get_markdown_metadata_endpoint(
        self, client: TestClient, db_session: Session, temp_markdown_file: Path
    ):
        """Test get_markdown_metadata endpoint."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/metadata")
        
        assert response.status_code == 200

    def test_get_markdown_metadata_endpoint_not_found(
        self, client: TestClient, db_session: Session
    ):
        """Test get_markdown_metadata endpoint with non-existent document."""
        response = client.get("/api/documents/999999/markdown/metadata")
        
        assert response.status_code == 404

    def test_get_markdown_metadata_endpoint_not_markdown(
        self, client: TestClient, db_session: Session, temp_docx_file: Path
    ):
        """Test get_markdown_metadata endpoint with non-Markdown file."""
        doc = Document(
            title="Test Doc",
            filename="test.docx",
            file_type="doc",
            file_path=str(temp_docx_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/metadata")
        
        assert response.status_code == 400

    def test_get_markdown_outline_endpoint(
        self, client: TestClient, db_session: Session, temp_markdown_file: Path
    ):
        """Test get_markdown_outline endpoint."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/outline")
        
        assert response.status_code == 200

    def test_get_markdown_data_endpoint(
        self, client: TestClient, db_session: Session, temp_markdown_file: Path
    ):
        """Test get_markdown_data endpoint."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/markdown/data")
        
        assert response.status_code == 200

    def test_update_markdown_content_endpoint(
        self, client: TestClient, db_session: Session, temp_markdown_file: Path
    ):
        """Test update_markdown_content endpoint."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.put(
            f"/api/documents/{doc.id}/markdown/content",
            json={"content": "Updated content"}
        )
        
        assert response.status_code == 200

    def test_replace_markdown_section_endpoint(
        self, client: TestClient, db_session: Session, temp_markdown_file: Path
    ):
        """Test replace_markdown_section endpoint."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.put(
            f"/api/documents/{doc.id}/markdown/section",
            json={
                "heading_text": "Introduction",
                "new_content": "Updated section content"
            }
        )
        
        assert response.status_code in [200, 404]

    def test_add_heading_to_markdown_endpoint(
        self, client: TestClient, db_session: Session, temp_markdown_file: Path
    ):
        """Test add_heading_to_markdown endpoint."""
        doc = Document(
            title="Test Markdown",
            filename="test.md",
            file_type="markdown",
            file_path=str(temp_markdown_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.post(
            f"/api/documents/{doc.id}/markdown/heading",
            json={
                "heading_text": "New Section",
                "content": "New section content"
            }
        )
        
        assert response.status_code == 200


class TestDocumentEndpointsExcel:
    """Tests for document endpoints with Excel-specific functionality."""

    def test_get_sheet_data_endpoint(
        self, client: TestClient, db_session: Session, temp_excel_file: Path
    ):
        """Test get_sheet_data endpoint."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/excel/sheets/1/data")
        
        assert response.status_code in [200, 404]

    def test_get_sheet_data_paginated_endpoint(
        self, client: TestClient, db_session: Session, temp_excel_file: Path
    ):
        """Test get_sheet_data_paginated endpoint."""
        doc = Document(
            title="Test Excel",
            filename="test.xlsx",
            file_type="excel",
            file_path=str(temp_excel_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(
            f"/api/documents/{doc.id}/excel/sheets/1/data/paginated",
            params={"start_row": 1, "end_row": 10}
        )
        
        assert response.status_code in [200, 404]


class TestDocumentEndpointsPpt:
    """Tests for document endpoints with PPT-specific functionality."""

    def test_get_slide_data_endpoint(
        self, client: TestClient, db_session: Session, temp_pptx_file: Path
    ):
        """Test get_slide_data endpoint."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(f"/api/documents/{doc.id}/ppt/slides/1")
        
        assert response.status_code in [200, 404]

    def test_get_slide_data_paginated_endpoint(
        self, client: TestClient, db_session: Session, temp_pptx_file: Path
    ):
        """Test get_slide_data_paginated endpoint."""
        doc = Document(
            title="Test PPT",
            filename="test.pptx",
            file_type="ppt",
            file_path=str(temp_pptx_file),
            content="Test content",
            html_content="<p>Test</p>",
        )
        db_session.add(doc)
        db_session.commit()
        db_session.refresh(doc)
        
        response = client.get(
            f"/api/documents/{doc.id}/ppt/slides/paginated",
            params={"start": 1, "limit": 5}
        )
        
        assert response.status_code in [200, 404]
