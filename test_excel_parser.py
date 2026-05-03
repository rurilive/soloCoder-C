import sys
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

sys.path.insert(0, str(Path(__file__).parent))

from app.parsers.excel_parser import (
    parse_excel,
    parse_excel_file,
    update_cell_in_excel,
    update_multiple_cells,
    add_new_sheet,
    delete_sheet,
    get_sheet_data_as_json,
)


def create_test_excel():
    """创建测试用的Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "销售数据"
    
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    headers = ["产品名称", "单价", "数量", "小计", "日期"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    data = [
        ["笔记本电脑", 5999, 10, "=B2*C2", datetime(2024, 1, 15)],
        ["无线鼠标", 99, 50, "=B3*C3", datetime(2024, 1, 16)],
        ["机械键盘", 399, 25, "=B4*C4", datetime(2024, 1, 17)],
        ["显示器", 1999, 8, "=B5*C5", datetime(2024, 1, 18)],
    ]
    
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in [2, 3, 4]:
                cell.alignment = Alignment(horizontal="right")
    
    ws.cell(row=7, column=1, value="总计")
    ws.cell(row=7, column=4, value="=SUM(D2:D5)")
    
    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 15
    
    ws2 = wb.create_sheet("库存统计")
    ws2.merge_cells('A1:C1')
    ws2['A1'] = "库存汇总表"
    ws2['A1'].font = Font(bold=True, size=14)
    ws2['A1'].alignment = Alignment(horizontal="center")
    
    ws2.append(["仓库", "产品", "数量"])
    ws2.append(["华东仓", "笔记本电脑", 100])
    ws2.append(["华东仓", "无线鼠标", 500])
    ws2.append(["华北仓", "机械键盘", 300])
    
    test_file = Path("test_excel.xlsx")
    wb.save(test_file)
    print(f"✅ 测试Excel文件已创建: {test_file}")
    return test_file


def test_parse_excel(file_path):
    """测试基础解析功能"""
    print("\n" + "=" * 60)
    print("测试1: 基础解析 (parse_excel)")
    print("=" * 60)
    
    text_content, html_content = parse_excel(str(file_path))
    
    print(f"文本内容长度: {len(text_content)} 字符")
    print(f"HTML内容长度: {len(html_content)} 字符")
    
    if text_content and "销售数据" in text_content:
        print("✅ 解析成功，包含预期内容")
    else:
        print("❌ 解析失败或内容不完整")
    
    return True


def test_parse_excel_file(file_path):
    """测试详细解析功能"""
    print("\n" + "=" * 60)
    print("测试2: 详细解析 (parse_excel_file)")
    print("=" * 60)
    
    result = parse_excel_file(str(file_path))
    
    if not result.success:
        print(f"❌ 解析失败: {result.error_message}")
        return False
    
    print(f"工作表数量: {len(result.sheets)}")
    for sheet in result.sheets:
        print(f"  - 工作表: {sheet.name}")
        print(f"    行数: {sheet.max_row}, 列数: {sheet.max_col}")
        print(f"    合并单元格: {len(sheet.merged_cells)} 个")
        
        if sheet.rows and sheet.rows[0]:
            first_row = sheet.rows[0]
            print(f"    第一行数据类型: {[c.data_type for c in first_row[:5]]}")
    
    print("✅ 详细解析成功")
    return True


def test_get_sheet_data_as_json(file_path):
    """测试JSON导出功能"""
    print("\n" + "=" * 60)
    print("测试3: JSON导出 (get_sheet_data_as_json)")
    print("=" * 60)
    
    data = get_sheet_data_as_json(str(file_path))
    
    if not data.get('success'):
        print(f"❌ 导出失败: {data.get('error')}")
        return False
    
    print(f"工作表数量: {data.get('sheet_count')}")
    for sheet in data.get('sheets', []):
        print(f"  - {sheet['name']}: {sheet['max_row']}行 x {sheet['max_col']}列")
    
    print("✅ JSON导出成功")
    return True


def test_update_cell(file_path):
    """测试单元格更新功能"""
    print("\n" + "=" * 60)
    print("测试4: 单元格更新 (update_cell_in_excel)")
    print("=" * 60)
    
    original_data = get_sheet_data_as_json(str(file_path))
    original_value = None
    for sheet in original_data.get('sheets', []):
        if sheet['name'] == '销售数据':
            if sheet['data'] and len(sheet['data']) >= 2:
                original_value = sheet['data'][1][0].get('value')
                break
    
    print(f"原始值 (A2): {original_value}")
    
    success = update_cell_in_excel(
        str(file_path),
        "销售数据",
        row=2,
        col=1,
        value="游戏笔记本电脑"
    )
    
    if not success:
        print("❌ 更新失败")
        return False
    
    new_data = get_sheet_data_as_json(str(file_path))
    new_value = None
    for sheet in new_data.get('sheets', []):
        if sheet['name'] == '销售数据':
            if sheet['data'] and len(sheet['data']) >= 2:
                new_value = sheet['data'][1][0].get('value')
                break
    
    print(f"更新后的值 (A2): {new_value}")
    
    if new_value == "游戏笔记本电脑":
        print("✅ 单元格更新成功")
        return True
    else:
        print("❌ 单元格值未正确更新")
        return False


def test_batch_update(file_path):
    """测试批量更新功能"""
    print("\n" + "=" * 60)
    print("测试5: 批量更新 (update_multiple_cells)")
    print("=" * 60)
    
    updates = [
        {"sheet_name": "销售数据", "row": 3, "col": 2, "value": 129},
        {"sheet_name": "销售数据", "row": 3, "col": 3, "value": 100},
    ]
    
    success = update_multiple_cells(str(file_path), updates)
    
    if success:
        print("✅ 批量更新成功")
        return True
    else:
        print("❌ 批量更新失败")
        return False


def test_add_and_delete_sheet(file_path):
    """测试添加和删除工作表"""
    print("\n" + "=" * 60)
    print("测试6: 工作表管理 (add_new_sheet, delete_sheet)")
    print("=" * 60)
    
    original_data = get_sheet_data_as_json(str(file_path))
    original_count = original_data.get('sheet_count', 0)
    print(f"原始工作表数量: {original_count}")
    
    success = add_new_sheet(str(file_path), "新工作表")
    if not success:
        print("❌ 添加工作表失败")
        return False
    
    after_add_data = get_sheet_data_as_json(str(file_path))
    after_add_count = after_add_data.get('sheet_count', 0)
    print(f"添加后工作表数量: {after_add_count}")
    
    if after_add_count != original_count + 1:
        print("❌ 工作表数量未增加")
        return False
    
    print("✅ 添加工作表成功")
    
    success = delete_sheet(str(file_path), "新工作表")
    if not success:
        print("❌ 删除工作表失败")
        return False
    
    after_delete_data = get_sheet_data_as_json(str(file_path))
    after_delete_count = after_delete_data.get('sheet_count', 0)
    print(f"删除后工作表数量: {after_delete_count}")
    
    if after_delete_count != original_count:
        print("❌ 工作表数量未恢复")
        return False
    
    print("✅ 删除工作表成功")
    return True


def main():
    print("=" * 60)
    print("Excel解析器和编辑器测试")
    print("=" * 60)
    
    test_file = create_test_excel()
    
    tests = [
        ("基础解析", test_parse_excel),
        ("详细解析", test_parse_excel_file),
        ("JSON导出", test_get_sheet_data_as_json),
        ("单元格更新", test_update_cell),
        ("批量更新", test_batch_update),
        ("工作表管理", test_add_and_delete_sheet),
    ]
    
    results = []
    for name, test_func in tests:
        try:
            success = test_func(test_file)
            results.append((name, success))
        except Exception as e:
            print(f"❌ 测试 {name} 抛出异常: {e}")
            import traceback
            traceback.print_exc()
            results.append((name, False))
    
    print("\n" + "=" * 60)
    print("测试结果汇总")
    print("=" * 60)
    
    passed = 0
    for name, success in results:
        status = "✅ 通过" if success else "❌ 失败"
        print(f"  {name}: {status}")
        if success:
            passed += 1
    
    print(f"\n总计: {passed}/{len(results)} 测试通过")
    
    if test_file.exists():
        print(f"\n测试文件: {test_file} (如需可手动删除)")
    
    return passed == len(results)


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
